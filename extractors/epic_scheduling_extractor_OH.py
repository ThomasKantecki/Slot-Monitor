"""Explore an Epic/MyChart open-scheduling flow and export public scheduling JSON.

This is intentionally an exploratory collector: it captures the site's own
XHR/fetch responses and visible UI metadata, without submitting patient data
or completing a booking.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import os
import re
import shutil
import sys
import threading
import time
import traceback
from http.cookiejar import CookieJar
from urllib.parse import parse_qs, urlencode, urlsplit
from urllib.request import HTTPCookieProcessor, Request, build_opener, urlopen
from pathlib import Path
from typing import Any


def use_shared_venv_if_needed():
    """Re-launch from the existing project venv when Run uses another Python."""
    try:
        import selenium  # noqa: F401
        return
    except ModuleNotFoundError as error:
        if error.name != "selenium":
            raise
    here = Path(__file__).resolve()
    candidates = [
        here.parent / ".venv" / "Scripts" / "python.exe",
        Path(r"C:\Users\PBA96B\ONEDRI~1\DOCUME~1\DATALI~1\.venv\Scripts\python.exe"),
    ]
    for interpreter in candidates:
        if interpreter.exists() and interpreter.resolve() != Path(sys.executable).resolve():
            os.execv(str(interpreter), [str(interpreter), str(here), *sys.argv[1:]])
    raise ModuleNotFoundError(
        "selenium is not installed in the selected Python interpreter. "
        "Select the project .venv interpreter or install selenium there."
    )


if "--guided-browser" in sys.argv:
    use_shared_venv_if_needed()
    from selenium import webdriver
    from selenium.webdriver.common.action_chains import ActionChains
    from selenium.common.exceptions import TimeoutException, WebDriverException
    from selenium.webdriver.chrome.options import Options as ChromeOptions
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.edge.options import Options as EdgeOptions


URL = "https://mychart.orlandohealth.com/MyChart/openscheduling"
SITE_CODE = "OH"
ACTIVE_RUN_TAG = ""
SPECIALTY_ALLOWLIST: set[str] | None = None
JSON_HINTS = ("api", "schedule", "appointment", "provider", "special", "location", "search")
DATE_RE = re.compile(r"\b(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)?\s*,?\s*(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}\b", re.I)


def extraction_folder_name():
    """Readable run folder: 24aug_OH_extr_144641 (24-hour local time)."""
    return time.strftime("%d%b").lower() + "_OH_extr_" + time.strftime("%H%M%S")


def artifact_name(filename: str, latest: bool = False) -> str:
    """Give every generated artifact an explicit OH/AH site identity."""
    if latest:
        return f"latest_{SITE_CODE}_{filename}"
    return f"{ACTIVE_RUN_TAG or SITE_CODE}_{filename}"


def set_active_run_tag(run_tag: str) -> None:
    """Use the readable dated run-folder name in every internal filename."""
    global ACTIVE_RUN_TAG
    ACTIVE_RUN_TAG = run_tag


def filter_allowed_specialties(specialties: list[dict[str, Any]], include_all: bool = False):
    """Apply an exact site-specific comparison allowlist when configured."""
    if include_all or not SPECIALTY_ALLOWLIST:
        return specialties
    allowed = {normalized_choice(name) for name in SPECIALTY_ALLOWLIST}
    return [item for item in specialties if normalized_choice(item.get("Name")) in allowed]

API_ENDPOINTS = {
    "bootstrap": "/MyChart/openscheduling",
    "workflow": "/MyChart/Scheduling/Anonymous/GetSchedulingWorkflowData",
    "specialty": "/MyChart/Scheduling/Anonymous/GetSpecialtyData",
    "decision_tree": "/MyChart/DecisionTrees/AnonymousDecisionTree/NextStep",
    "questionnaire_evaluation": "/MyChart/Scheduling/Anonymous/EvaluateQuestionnaireAnswers",
    "slots": "/MyChart/Scheduling/Anonymous/GetSlots",
}


class OrlandoHealthApiSession:
    """Small stdlib-only client for the public anonymous OpenScheduling APIs."""

    def __init__(self):
        self.base_url = "https://mychart.orlandohealth.com"
        self.opener = build_opener(HTTPCookieProcessor(CookieJar()))
        self.csrf_token = ""
        self.page_nonce = ""
        self.captures: list[dict[str, Any]] = []

    def bootstrap(self):
        url = self.base_url + API_ENDPOINTS["bootstrap"]
        request = Request(url, headers={"Accept": "text/html,application/xhtml+xml", "User-Agent": "Mozilla/5.0"})
        with self.opener.open(request, timeout=45) as response:
            body = response.read().decode("utf-8", errors="replace")
            status = response.status
            mime = response.headers.get("content-type", "")
        self.csrf_token = (
            re.search(r'name=["\']__RequestVerificationToken["\'][^>]*value=["\']([^"\']+)', body, re.I)
            or re.search(r'value=["\']([^"\']+)["\'][^>]*name=["\']__RequestVerificationToken', body, re.I)
        )
        self.csrf_token = self.csrf_token.group(1) if self.csrf_token else ""
        nonce_match = re.search(r'<script[^>]+id=["\']cspScripts["\'][^>]*nonce=["\']([^"\']*)', body, re.I)
        self.page_nonce = nonce_match.group(1) if nonce_match else ""
        self.captures.append({
            "source": "api_inventory",
            "method": "GET",
            "url": url,
            "status": status,
            "mime_type": mime,
            "body": body,
            "data": None,
        })
        if not self.csrf_token:
            raise RuntimeError("Could not find the anonymous scheduling CSRF token in the bootstrap page.")

    def post_json(self, endpoint_key: str, fields: Any, extra: dict[str, Any] | None = None):
        path = API_ENDPOINTS[endpoint_key]
        url = self.base_url + path + "?noCache=" + str(time.time_ns())
        if isinstance(fields, dict):
            pairs = [(key, "" if value is None else str(value)) for key, value in fields.items()]
        else:
            pairs = [(key, "" if value is None else str(value)) for key, value in fields]
        encoded = urlencode(pairs)
        headers = {
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "X-Requested-With": "XMLHttpRequest",
            "Origin": self.base_url,
            "Referer": self.base_url + API_ENDPOINTS["bootstrap"],
            "User-Agent": "Mozilla/5.0",
        }
        request = Request(url, data=encoded.encode("utf-8"), headers=headers, method="POST")
        try:
            with self.opener.open(request, timeout=45) as response:
                body = response.read().decode("utf-8", errors="replace")
                status = response.status
                mime = response.headers.get("content-type", "")
        except Exception as error:
            self.captures.append({"source": "api_inventory", "method": "POST", "url": url, "post_data": encoded, "status": 0, "body_error": str(error), "data": None, **(extra or {})})
            raise
        try:
            data = json.loads(body)
        except (TypeError, ValueError):
            data = None
        capture = {
            "source": "api_inventory",
            "method": "POST",
            "url": url,
            "post_data": encoded,
            "status": status,
            "mime_type": mime,
            "body": body,
            "data": data,
        }
        if extra:
            capture.update(extra)
        self.captures.append(capture)
        if not isinstance(data, dict):
            raise RuntimeError(f"{path} returned non-JSON data (HTTP {status}).")
        return data


def run_api_inventory(args):
    """Map and capture the public API-backed inventory without opening Edge."""
    session_name = extraction_folder_name()
    set_active_run_tag(session_name)
    artifact_dir = Path(args.artifacts) / session_name
    artifact_dir.mkdir(parents=True, exist_ok=True)
    client = OrlandoHealthApiSession()
    try:
        client.bootstrap()
        workflow = client.post_json("workflow", {
            "schedulingParameters.isAnonymous": "true",
            "schedulingParameters.workflow": "NewProvider",
            "nonce": client.page_nonce,
            "__RequestVerificationToken": client.csrf_token,
        }, {"phase": "workflow"})
        specialties = filter_allowed_specialties(workflow.get("Specialties", []), args.all_specialties)
        if args.limit_specialties:
            specialties = specialties[:args.limit_specialties]
        print(f"API-first inventory mapping {len(specialties)} specialties.", flush=True)
        for index, specialty in enumerate(specialties, start=1):
            name = specialty.get("Name", "")
            print(f"  {index}/{len(specialties)} {name}", flush=True)
            detail = client.post_json("specialty", {
                "SpecialtyId": specialty.get("Id", specialty.get("ID", "")),
                "isFirstLoad": "true",
                "schedulingOverridesString": "{}",
                "__RequestVerificationToken": client.csrf_token,
            }, {"phase": "specialty", "specialty_id": specialty.get("Id", specialty.get("ID", "")), "specialty_name": name})
            detail["_captured_specialty"] = {"id": specialty.get("Id", specialty.get("ID", "")), "name": name}
        capture_path = artifact_dir / artifact_name("network_capture.json")
        capture_path.write_text(json.dumps(client.captures, indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("api_endpoint_map.json")).write_text(json.dumps({
            "observed": build_request_index(client.captures),
            "known_public_workflow_endpoints": API_ENDPOINTS,
            "notes": [
                "Workflow and specialty detail were captured directly with a cookie-backed anonymous session.",
                "Decision-tree and GetSlots endpoints are mapped from the validated browser trace and are the next API-first implementation phase.",
                "No login, hold, reservation, confirmation, or booking endpoint is used."
            ],
        }, indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("workflow_catalog.json")).write_text(json.dumps(build_workflow_catalog(client.captures), indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("entity_catalog.json")).write_text(json.dumps(build_entity_catalog(client.captures), indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("api_inventory_summary.json")).write_text(json.dumps({
            "specialty_count": len(specialties),
            "specialties": [{
                "name": item.get("specialty_name", ""),
                "provider_count": len((item.get("data") or {}).get("Providers", [])),
                "department_count": len((item.get("data") or {}).get("Departments", [])),
                "location_count": len((item.get("data") or {}).get("Locations", [])),
                "provider_department_pair_count": len((item.get("data") or {}).get("ProviderDepartmentPairs", [])),
                "visit_type_count": len((item.get("data") or {}).get("VisitTypes", [])),
                "reason_for_visit_count": len((item.get("data") or {}).get("ReasonsForVisit", [])),
            } for item in client.captures if item.get("phase") == "specialty"
            ],
            "captured_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        }, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Captured {len(client.captures)} API responses in {artifact_dir}")
    except Exception:
        (artifact_dir / artifact_name("error.txt")).write_text(traceback.format_exc(), encoding="utf-8")
        print(f"API-first mapping stopped. Details saved to {artifact_dir / artifact_name('error.txt')}")


def epic_id(item: dict[str, Any]) -> Any:
    return item.get("ID", item.get("Id", ""))


def form_scalar(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def postify(value: Any, prefix: str = "", output: list[tuple[str, str]] | None = None):
    """Flatten Epic's nested view models into ASP.NET form-field names."""
    if output is None:
        output = []
    if isinstance(value, dict):
        for key, child in value.items():
            name = f"{prefix}.{key}" if prefix else str(key)
            postify(child, name, output)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            postify(child, f"{prefix}[{index}]", output)
    elif value is not None and prefix:
        output.append((prefix, form_scalar(value)))
    return output


def raw_workflow(settings: dict[str, Any]):
    return {
        "Type": settings.get("WorkflowType"),
        "FinderType": settings.get("FinderType"),
        "IsGuest": False,
        "IsAnonymous": True,
        "IsFromPrelogin": settings.get("IsFromPrelogin"),
        "RescheduleDat": None,
        "RootDecisionTreeId": settings.get("RootDecisionTreeId"),
        "DecisionTreeAnswerId": settings.get("DecisionTreeAnswerId"),
        "DecisionTreeNodeId": settings.get("DecisionTreeNodeId"),
        "DecisionTreeNodeCsn": settings.get("DecisionTreeNodeCsn"),
        "SchedulingControllerParams": {},
        "SecureSessionToken": None,
        "BrowserId": None,
        "MultiApptSlotStepStartingView": None,
        "MultiApptSlotStepSelectionView": None,
        "IsAuthenticatedWidget": False,
    }


def normalized_choice(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def usable_question_choices(question: dict[str, Any]):
    return [str(choice.get("Text", "")).strip() for choice in question.get("Choices", [])
            if normalized_choice(choice.get("Text")) not in {"", "choose", "[choose]"}]


def match_choice(question: dict[str, Any], wanted: str):
    target = normalized_choice(wanted)
    choices = question.get("Choices", [])
    return (
        next((choice for choice in choices if normalized_choice(choice.get("Text")) == target), None)
        or next((choice for choice in choices if normalized_choice(choice.get("Text")).startswith(target)), None)
        or next((choice for choice in choices if target in normalized_choice(choice.get("Text"))), None)
    )


def select_default_decision_answer(question: dict[str, Any], visit_type: dict[str, Any]):
    prompt = normalized_choice(question.get("Prompt"))
    choices = usable_question_choices(question)
    visit_label = normalized_choice(visit_type.get("DisplayName") or visit_type.get("Name"))
    if "name of your pcp" in prompt or "ordering provider" in prompt:
        return "None"
    if "last period" in prompt or "menstruation" in prompt:
        return "06/01/2026"
    if "orlando health" in prompt and ("florida medical" in prompt or "fhv" in prompt):
        preferred = "Florida Medical Clinic Orlando Health" if "florida medical clinic" in visit_label else "Orlando Health"
        return preferred if match_choice(question, preferred) else choices[0]
    if "new patient" in prompt or "longer than three years" in prompt or "longer than 3 years" in prompt:
        return "Yes" if match_choice(question, "Yes") else choices[0]
    if "vaccinated" in prompt or "due for a well child" in prompt:
        return "Yes" if match_choice(question, "Yes") else choices[0]
    if "established with orlando health physician associates" in prompt:
        return "Yes" if match_choice(question, "Yes") else choices[0]
    if any(text in prompt for text in ("have you seen", "been seen", "established patient", "existing patient")):
        return "No" if match_choice(question, "No") else choices[0]
    if "medical emergency" in prompt or "immediate medical attention" in prompt or "call 911" in prompt:
        return "No" if match_choice(question, "No") else choices[0]
    if "primary insurance" in prompt and ("medicaid" in prompt or "careplus" in prompt):
        return "No" if match_choice(question, "No") else choices[0]
    if "type of insurance" in prompt:
        for answer in ("PPO", "Other/None", "Other"):
            if match_choice(question, answer):
                return answer
    if "following coverages" in prompt or ("coverage" in prompt and match_choice(question, "Other/None")):
        return "Other/None" if match_choice(question, "Other/None") else choices[-1]
    if "pregnancy" in prompt and "reason" in prompt:
        wanted = "Yes" if "new ob patient" in visit_label else "No"
        return wanted if match_choice(question, wanted) else choices[0]
    if "prior surgery" in prompt:
        return "No" if match_choice(question, "No") else choices[0]
    if "car accident" in prompt or "workers compensation" in prompt or "worker's compensation" in prompt:
        return "No" if match_choice(question, "No") else choices[0]
    if "age" in prompt:
        for answer in ("18 years old and older", "18 years or older", "18 or older", "Yes"):
            if match_choice(question, answer):
                return answer
        if match_choice(question, "30"):
            return "30"
    if "medicaid" in prompt or "careplus" in prompt:
        if match_choice(question, "Not Medicaid"):
            return "Not Medicaid"
        if match_choice(question, "No"):
            return "No"
    if "reason for visit" in prompt:
        for answer in ("Check up", "New Patient", "Physical", "Lab"):
            if match_choice(question, answer):
                return answer
        return next((choice for choice in choices if normalized_choice(choice) != "other"), choices[0] if choices else "")
    if "what kind of appointment" in prompt:
        for answer in ("New Patient Visit", "Annual Physical", "Sick Visit"):
            if match_choice(question, answer):
                return answer
    if "body part" in prompt:
        for answer in ("Knee", "Knee Pain", "Shoulder", "Arm Pain", "Not listed"):
            if match_choice(question, answer):
                return answer
    if "following diagnoses" in prompt or "diagnosis" in prompt:
        for answer in ("Abdominal Pain", "Knee Pain", "Other", "Not Sure"):
            if match_choice(question, answer):
                return answer
    if "which side" in prompt:
        return "Right" if match_choice(question, "Right") else choices[0]
    normalized_choices = {normalized_choice(choice) for choice in choices}
    if normalized_choices == {"yes", "no"}:
        return "No"
    return choices[0] if choices else ""


def core_question_answer(question: dict[str, Any], answer_text: str):
    choice = match_choice(question, answer_text)
    if not choice and question.get("Choices"):
        raise RuntimeError(f"Could not answer decision-tree prompt {question.get('Prompt')!r} with {answer_text!r}.")
    return {
        "ID": question.get("ID"),
        "DAT": question.get("DAT"),
        "QuestionType": question.get("QuestionType"),
        "ResponseType": question.get("ResponseType"),
        "IsRequired": question.get("IsRequired"),
        "IsMultiResponse": question.get("IsMultiResponse"),
        "IsTrigger": question.get("IsTrigger"),
        "IsEnabled": question.get("IsEnabled"),
        "DisplayStyle": question.get("DisplayStyle"),
        "DisplayStyleVal": question.get("DisplayStyleVal"),
        "Answer": {"Choices": [{"Index": choice.get("Index")}]} if choice else answer_text,
    }


def traverse_default_decision_tree(client: OrlandoHealthApiSession, visit_type: dict[str, Any], workflow: dict[str, Any]):
    tree_id = visit_type.get("AnonymousSchedulingDecisionTreeId")
    if not tree_id:
        return {"prompts": [], "override": None, "evaluated": {}, "stop": False, "message": ""}
    traversal = {
        "TreeID": tree_id,
        "TreeAnswerID": None,
        "IsTraversalComplete": False,
        "SourceWorkflow": 5,
        "TreeWasDirty": False,
        "TreeWasLocked": False,
        "RestartTree": True,
        "UseInProgress": False,
        "AdditionalContext": {
            "VisitTypeID": epic_id(visit_type),
            "TicketID": "",
            "AppointmentRequestIds": [],
            "OriginalApptDAT": "",
            "FavoriteApptDAT": "",
            "OrdersString": "",
            "IsGuest": False,
            "SchedulingWorkflowType": workflow.get("Type"),
            "TermIds": [],
            "SchedGrouperIds": None,
            "IsAuthenticatedWidget": False,
        },
        "ParentAnswerID": None,
    }
    fields = postify({"traversalInfo": traversal})
    fields.append(("__RequestVerificationToken", client.csrf_token))
    response = client.post_json("decision_tree", fields, {"phase": "decision_tree_start"})
    prompts = []
    while not (response.get("TraversalInfo") or {}).get("IsTraversalComplete"):
        node = response.get("NextInputNode") or {}
        question = node.get("Question") or {}
        if not question:
            raise RuntimeError("Decision tree returned no question before completion.")
        answer = select_default_decision_answer(question, visit_type)
        prompts.append({"prompt": question.get("Prompt", ""), "answer": answer,
                        "choices": usable_question_choices(question)})
        traversal = json.loads(json.dumps(response.get("TraversalInfo") or {}))
        traversal["RestartTree"] = False
        payload = {
            "traversalInfo": traversal,
            "prevInputNode": {
                "CSN": node.get("CSN"), "ID": node.get("ID"), "Type": node.get("Type"),
                "IsFirst": node.get("IsFirst"), "Question": None, "Questionnaire": None,
                "DecisionTree": None, "DeclutterNavigationButtons": node.get("DeclutterNavigationButtons"),
            },
            "question": core_question_answer(question, answer),
        }
        fields = postify(payload)
        fields.append(("__RequestVerificationToken", client.csrf_token))
        response = client.post_json("decision_tree", fields, {
            "phase": "decision_tree_answer", "question_prompt": question.get("Prompt", ""), "answer": answer,
        })
        if len(prompts) > 15:
            raise RuntimeError("Decision tree exceeded the expected question count.")
    return {"prompts": prompts, "tree_answer_id": (response.get("TraversalInfo") or {}).get("TreeAnswerID")}


def evaluate_questionnaire(client: OrlandoHealthApiSession, workflow: dict[str, Any], specialty: dict[str, Any],
                           reason: dict[str, Any], visit_type: dict[str, Any]):
    traversal = traverse_default_decision_tree(client, visit_type, workflow)
    tree_id = visit_type.get("AnonymousSchedulingDecisionTreeId")
    if not tree_id:
        return {**traversal, "override": None, "evaluated": {}, "stop": False, "message": ""}
    override = {
        "LqfIds": [tree_id],
        "HqaIds": [traversal.get("tree_answer_id")],
        "OriginalPrcId": epic_id(visit_type),
        "OriginalRfv": reason.get("CategoryValue"),
        "OriginalRfvLine": reason.get("Id", reason.get("ID")),
    }
    fields = postify({
        "workflow": workflow,
        "schedulingOverridesString": json.dumps(override, separators=(",", ":")),
        "termIds": [],
        "nonce": client.page_nonce,
    })
    fields.append(("__RequestVerificationToken", client.csrf_token))
    evaluated = client.post_json("questionnaire_evaluation", fields, {"phase": "questionnaire_evaluation"})
    message = re.sub(r"<[^>]+>", " ", str(evaluated.get("Instructions") or ""))
    return {**traversal, "override": override, "evaluated": evaluated,
            "stop": bool(evaluated.get("StopScheduling")), "message": " ".join(message.split())}


def choose_slot_reason(reasons: list[dict[str, Any]], visit_type: dict[str, Any]):
    candidates = [reason for reason in reasons if reason.get("CanDirectSchedule") is not False] or reasons
    visit_label = normalized_choice(visit_type.get("DisplayName") or visit_type.get("Name"))
    # Exact names matter here: several specialties expose both Orlando Health
    # and Florida Medical Clinic reasons containing the words "new patient".
    exact = next((reason for reason in candidates if visit_label and visit_label in {
        normalized_choice(reason.get("Title")), normalized_choice(reason.get("DisplayName"))
    }), None)
    if exact:
        return exact
    visit_id = epic_id(visit_type)
    visit_match = next((reason for reason in candidates if visit_id and visit_id in {
        reason.get("DefaultVisitTypeId"), reason.get("VisitTypeId")
    }), None)
    if visit_match:
        return visit_match
    return next((reason for reason in candidates if any(
        label and (label == visit_label or label in visit_label or visit_label in label)
        for label in [normalized_choice(reason.get("Title")), normalized_choice(reason.get("DisplayName"))]
    )), candidates[0] if candidates else None)


def supported_visit_types(detail: dict[str, Any]):
    supported = {info.get("VisitTypeID") for pair in detail.get("ProviderDepartmentPairs", [])
                 for info in pair.get("VisitTypeInformation", [])}
    all_visits = detail.get("VisitTypes", [])
    visits = [visit for visit in all_visits if epic_id(visit) in supported] or all_visits
    new_patient = [visit for visit in visits if "new patient" in normalized_choice(
        visit.get("DisplayName") or visit.get("Name"))]
    # Process every public new-patient visit type. This commonly includes a
    # separate Orlando Health and Florida Medical Clinic path.
    return new_patient or visits


def build_slot_request(workflow_settings: dict[str, Any], workflow: dict[str, Any], specialty: dict[str, Any],
                       reason: dict[str, Any], visit_type: dict[str, Any], pairs: list[dict[str, Any]], lqf: dict[str, Any]):
    override = lqf.get("override") or {}
    return {
        "workflow": workflow,
        "appointmentBuilder": {
            "Appointments": [{
                "VisitTypeId": epic_id(visit_type), "PanelId": None, "BundleId": None,
                "RescheduleDat": None, "TicketId": None, "AppointmentRequestIds": [],
                "ProviderDepartmentPairs": [{
                    "ProviderId": pair.get("ProviderId"), "DepartmentId": pair.get("DepartmentId"),
                    "ChildProviderIds": pair.get("ChildProviderIds", []), "IsTeamMember": pair.get("IsTeamMember"),
                    "PoolLine": pair.get("PoolLine"), "PoolTier": pair.get("PoolTier"),
                } for pair in pairs],
                "Slot": "", "OrderIds": [], "LqfIds": override.get("LqfIds", []),
                "PatientAnswerIds": override.get("HqaIds", []),
                "OriginalVisitTypeId": override.get("OriginalPrcId") if override else None,
                "AddToWaitList": None, "SearchStartDte": None,
                "SelectedTelehealthMode": visit_type.get("DefaultTelehealthMode"),
                "HasTelehealthToggles": None, "HasUsedTelehealthToggles": None,
                "AccessCode": None, "CanSkipLicensureCheck": None,
                "InitialPoolLine": visit_type.get("PoolLine"), "MaxPoolLine": visit_type.get("MaxPoolLine"),
                "MaxPoolTier": visit_type.get("MaxPoolTier"),
            }],
            "ReasonForVisitLine": reason.get("LineInWDF15000") or reason.get("Id", reason.get("ID")),
            "ReasonForVisitValue": reason.get("CategoryValue"),
            "ReasonForVisitAllowProviderSelection": reason.get("AllowProviderSelect"),
            "UseInsuranceForVisit": "", "SpecialtyId": epic_id(specialty),
            "ClientIANATimeZone": "America/New_York", "SearchPriority": 0,
        },
        "startDte": workflow_settings.get("CurrentDTE"),
        "useSchedulingPreferences": False,
        "continueInfo": None,
    }


def parse_department_address(department: dict[str, Any]):
    address = department.get("Address", "")
    if isinstance(address, list):
        address = ", ".join(str(part) for part in address if part)
    address = str(address or "")
    match = re.search(r",\s*([^,]+?)\s+([A-Z]{2})\s+(\d{5})(?:-\d{4})?\s*$", address)
    return address, (match.group(1) if match else ""), (match.group(2) if match else ""), (match.group(3) if match else "")


def normalize_api_slot(slot: dict[str, Any], specialty: dict[str, Any], visit_type: dict[str, Any],
                       reason: dict[str, Any], prompts: list[dict[str, Any]], providers: dict[Any, dict[str, Any]],
                       departments: dict[Any, dict[str, Any]], load_number: int):
    provider = providers.get(slot.get("ProviderId"), {})
    department = departments.get(slot.get("DepartmentId"), {})
    address, city, state, zip_code = parse_department_address(department)
    return {
        "specialty": specialty.get("Name", ""),
        "appointment_type": visit_type.get("DisplayName") or visit_type.get("Name", ""),
        "reason_for_visit": reason.get("DisplayName") or reason.get("Title", ""),
        "provider_name": provider.get("Name", ""),
        "provider_id": provider.get("ID", slot.get("ProviderId", "")),
        "provider_credentials": provider.get("Credentials", ""),
        "location_name": department.get("Name", ""),
        "department_id": department.get("ID", slot.get("DepartmentId", "")),
        "address": address, "city": city, "state": state, "zip": zip_code,
        "appointment_date": slot.get("DateString", ""), "appointment_time": slot.get("TimeString", ""),
        "display_datetime_utc": slot.get("DisplayDateTimeUtc", ""), "days_ahead": slot.get("DaysAhead", ""),
        "length_minutes": slot.get("LengthInMinutes", ""), "timezone": slot.get("TimeZoneMarker", ""),
        "decision_tree_path": " | ".join(f"{item['prompt']} => {item['answer']}" for item in prompts),
        "load_number": load_number, "source_url": URL,
    }


def write_api_slot_artifacts(artifact_dir: Path, captures: list[dict[str, Any]], rows: list[dict[str, Any]],
                             summaries: list[dict[str, Any]], inventory_rows: list[dict[str, Any]],
                             publish_latest: bool = True):
    unique_rows = list({(
        row["specialty"], row["appointment_type"], row["provider_id"], row["department_id"],
        row["display_datetime_utc"], row["appointment_date"], row["appointment_time"]
    ): row for row in rows}.values())
    unique_inventory = list({(
        row["specialty"], row["provider_id"], row["department_id"]
    ): row for row in inventory_rows}.values())
    walk_in_rows = extract_walk_in_locations(captures)
    (artifact_dir / artifact_name("network_capture.json")).write_text(json.dumps(captures, indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / artifact_name("slots.json")).write_text(json.dumps(unique_rows, indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / artifact_name("flow_summary.json")).write_text(json.dumps(summaries, indent=2, ensure_ascii=False), encoding="utf-8")
    # Keep OH and AH appointment CSVs structurally identical even when a run
    # returns zero slots (for example, AdventHealth Lab endpoint errors).
    slot_fields = [
        "specialty", "appointment_type", "reason_for_visit", "provider_name", "provider_id",
        "provider_credentials", "location_name", "department_id", "address", "city", "state", "zip",
        "appointment_date", "appointment_time", "display_datetime_utc", "days_ahead", "length_minutes",
        "timezone", "decision_tree_path", "load_number", "source_url",
    ]
    inventory_fields = ["specialty", "provider_name", "provider_id", "provider_credentials", "location_name",
                        "department_id", "address", "city", "state", "zip", "supported_appointment_types"]
    walk_in_fields = ["clinic_name", "address", "city", "state", "zip", "hours_url", "source_url"]
    outputs = [
        (artifact_dir / artifact_name("slots.csv"), unique_rows, slot_fields),
        (artifact_dir / artifact_name("doctor_locations.csv"), unique_inventory, inventory_fields),
        (artifact_dir / artifact_name("walk_in_locations.csv"), walk_in_rows, walk_in_fields),
    ]
    if publish_latest:
        outputs.extend([
            (artifact_dir.parent / artifact_name("slots.csv", latest=True), unique_rows, slot_fields),
            (artifact_dir.parent / artifact_name("doctor_locations.csv", latest=True), unique_inventory, inventory_fields),
            (artifact_dir.parent / artifact_name("walk_in_locations.csv", latest=True), walk_in_rows, walk_in_fields),
        ])
        (artifact_dir.parent / artifact_name("flow_summary.json", latest=True)).write_text(
            json.dumps(summaries, indent=2, ensure_ascii=False), encoding="utf-8")
    for path, data, fields in outputs:
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(data)
    create_master_workbook(artifact_dir, publish_latest=publish_latest)
    return unique_rows, unique_inventory, walk_in_rows


def create_master_workbook(artifact_dir: Path, publish_latest: bool = True):
    """Create a filtered, readable workbook with one worksheet per CSV dataset."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    first = True
    sheet_specs = [
        (f"{SITE_CODE} Appointments", artifact_dir / artifact_name("slots.csv"), f"{SITE_CODE}AppointmentsTable"),
        (f"{SITE_CODE} Doctor Locations", artifact_dir / artifact_name("doctor_locations.csv"), f"{SITE_CODE}DoctorLocationsTable"),
        (f"{SITE_CODE} Walk-In Clinics", artifact_dir / artifact_name("walk_in_locations.csv"), f"{SITE_CODE}WalkInClinicsTable"),
    ]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for sheet_name, csv_path, table_name in sheet_specs:
        if not csv_path.exists():
            continue
        if first:
            sheet = workbook.active
            sheet.title = sheet_name
            first = False
        else:
            sheet = workbook.create_sheet(sheet_name)
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            data = list(reader)
        if not data:
            continue
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for row in data:
            sheet.append(row)
        max_row = len(data)
        max_col = len(data[0])
        header = sheet[1]
        for cell in header:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = cell_border
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = cell_border
        sheet.auto_filter.ref = f"A1:{sheet.cell(max_row, max_col).coordinate}"
        if max_row >= 2:
            table = Table(displayName=table_name, ref=sheet.auto_filter.ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                                  showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        for col_index, cell in enumerate(header, start=1):
            header_length = len(str(cell.value or "")) + 4
            content_length = max((len(str(sheet.cell(row, col_index).value or "")) for row in range(2, max_row + 1)), default=0)
            sheet.column_dimensions[cell.column_letter].width = min(max(header_length, min(content_length + 2, 60)), 60)

    if first:
        sheet = workbook.active
        sheet.title = "No Data"
        sheet["A1"] = "No CSV datasets were available for this run."
    output_path = artifact_dir / artifact_name("master.xlsx")
    workbook.save(output_path)
    if publish_latest:
        shutil.copy2(output_path, artifact_dir.parent / artifact_name("master.xlsx", latest=True))
    return output_path


def extract_walk_in_locations(captures: list[dict[str, Any]]):
    """Extract informational walk-in clinic locations from public instructions."""
    rows = []
    seen = set()
    for capture in captures:
        data = capture.get("data") or {}
        instructions = data.get("Instructions", "") if isinstance(data, dict) else ""
        if "walk-in clinic" not in html.unescape(str(instructions)).lower():
            continue
        text = html.unescape(re.sub(r"<style[\s\S]*?</style>|<[^>]+>", " ", str(instructions), flags=re.I))
        text = " ".join(text.split())
        hours_match = re.search(r'href=["\']([^"\']*walk-in-clinic-locations[^"\']*)', str(instructions), re.I)
        hours_url = hours_match.group(1) if hours_match else ""
        for match in re.finditer(r"([A-Z][A-Za-z'’ .&-]{2,60}?)\s*\(([^()]+?,\s*FL\s+\d{5}(?:-\d{4})?)\)", text):
            name = " ".join(match.group(1).split()).strip(" -")
            address = " ".join(match.group(2).split())
            parsed = re.search(r"(.+?)\s+([A-Za-z .'-]+),\s*(FL)\s+(\d{5}(?:-\d{4})?)$", address)
            city = parsed.group(2).strip() if parsed else ""
            state = parsed.group(3) if parsed else "FL"
            zip_code = parsed.group(4) if parsed else ""
            key = (name, address)
            if key in seen:
                continue
            seen.add(key)
            rows.append({"clinic_name": name, "address": address, "city": city, "state": state,
                         "zip": zip_code, "hours_url": hours_url, "source_url": capture.get("url", "")})
        # These two official entries contain nested parentheses or omit the
        # comma before the state, so they need explicit tolerant matching.
        for name, address in (
            ("Windermere (Summerport)", "5151 Windermere (Summerport) Vineland Rd, Ste. 206, Windermere, FL 34786"),
            ("Melbourne", "5565 N. Wickham Rd. Melbourne FL 32940"),
        ):
            if name.lower() not in text.lower() or address.lower() not in text.lower():
                continue
            key = (name, address)
            if key in seen:
                continue
            seen.add(key)
            parsed = re.search(r"(.+?)\s+([A-Za-z .'-]+),?\s*(FL)\s+(\d{5}(?:-\d{4})?)$", address)
            rows.append({"clinic_name": name, "address": address,
                         "city": parsed.group(2).strip() if parsed else "",
                         "state": "FL", "zip": parsed.group(4) if parsed else "",
                         "hours_url": hours_url, "source_url": capture.get("url", "")})
    return rows


def run_api_slots(args):
    """Directly traverse the public workflow and capture paginated slots."""
    session_name = extraction_folder_name()
    set_active_run_tag(session_name)
    artifact_dir = Path(args.artifacts) / session_name
    artifact_dir.mkdir(parents=True, exist_ok=True)
    client = OrlandoHealthApiSession()
    rows: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    inventory_rows: list[dict[str, Any]] = []
    try:
        client.bootstrap()
        workflow_data = client.post_json("workflow", {
            "schedulingParameters.isAnonymous": "true", "schedulingParameters.workflow": "NewProvider",
            "nonce": client.page_nonce, "__RequestVerificationToken": client.csrf_token,
        }, {"phase": "workflow"})
        workflow_settings = workflow_data.get("WorkflowSettings", {})
        workflow = raw_workflow(workflow_settings)
        specialties = filter_allowed_specialties(workflow_data.get("Specialties", []), args.all_specialties)
        if args.specialty:
            wanted = normalized_choice(args.specialty)
            specialties = [item for item in specialties if wanted in normalized_choice(item.get("Name"))]
        if args.limit_specialties:
            specialties = specialties[:args.limit_specialties]
        print(f"Direct slot extraction starting for {len(specialties)} specialties.", flush=True)
        for specialty_index, specialty in enumerate(specialties, start=1):
            name = specialty.get("Name", "")
            try:
                base_detail = client.post_json("specialty", {
                    "SpecialtyId": epic_id(specialty), "isFirstLoad": "true", "schedulingOverridesString": "{}",
                    "__RequestVerificationToken": client.csrf_token,
                }, {"phase": "specialty", "specialty_name": name})
            except Exception as specialty_error:
                summaries.append({
                    "specialty": name, "appointment_type": "", "status": "specialty_response_error",
                    "message": str(specialty_error), "slot_count": 0,
                })
                print(f"  {specialty_index}/{len(specialties)} {name}: skipped ({specialty_error}); continuing", flush=True)
                continue
            base_providers = {item.get("ID"): item for item in base_detail.get("Providers", [])}
            base_departments = {item.get("ID"): item for item in base_detail.get("Departments", [])}
            for pair in base_detail.get("ProviderDepartmentPairs", []):
                provider = base_providers.get(pair.get("ProviderId"), {})
                department = base_departments.get(pair.get("DepartmentId"), {})
                address, city, state, zip_code = parse_department_address(department)
                visit_names = []
                visit_ids = {info.get("VisitTypeID") for info in pair.get("VisitTypeInformation", [])}
                for item in base_detail.get("VisitTypes", []):
                    if epic_id(item) in visit_ids:
                        visit_names.append(item.get("DisplayName") or item.get("Name", ""))
                inventory_rows.append({
                    "specialty": name, "provider_name": provider.get("Name", ""),
                    "provider_id": provider.get("ID", pair.get("ProviderId", "")),
                    "provider_credentials": provider.get("Credentials", ""),
                    "location_name": department.get("Name", ""),
                    "department_id": department.get("ID", pair.get("DepartmentId", "")),
                    "address": address, "city": city, "state": state, "zip": zip_code,
                    "supported_appointment_types": " | ".join(visit_names),
                })
            visits = supported_visit_types(base_detail)
            print(f"  {specialty_index}/{len(specialties)} {name}: {len(visits)} visit flow(s)", flush=True)
            for visit_type in visits:
                reason = choose_slot_reason(base_detail.get("ReasonsForVisit", []), visit_type)
                if not reason:
                    summaries.append({"specialty": name, "appointment_type": visit_type.get("DisplayName", ""),
                                      "status": "no_reason_for_visit", "slot_count": 0})
                    continue
                try:
                    lqf = evaluate_questionnaire(client, workflow, specialty, reason, visit_type)
                except Exception as flow_error:
                    summaries.append({
                        "specialty": name,
                        "appointment_type": visit_type.get("DisplayName") or visit_type.get("Name", ""),
                        "status": "flow_error", "message": str(flow_error), "slot_count": 0,
                    })
                    print(f"    {summaries[-1]['appointment_type']}: skipped ({flow_error})", flush=True)
                    continue
                if lqf.get("stop"):
                    summaries.append({"specialty": name, "appointment_type": visit_type.get("DisplayName", ""),
                                      "status": "public_stop", "message": lqf.get("message", ""), "slot_count": 0,
                                      "decision_tree": lqf.get("prompts", [])})
                    continue
                detail = base_detail
                evaluated = lqf.get("evaluated") or {}
                active_visit = visit_type
                if lqf.get("override"):
                    try:
                        detail = client.post_json("specialty", {
                            "SpecialtyId": epic_id(specialty), "isFirstLoad": "false",
                            "schedulingOverridesString": json.dumps(lqf["override"], separators=(",", ":")),
                            "__RequestVerificationToken": client.csrf_token,
                        }, {"phase": "specialty_override", "specialty_name": name})
                    except Exception as override_error:
                        summaries.append({
                            "specialty": name,
                            "appointment_type": visit_type.get("DisplayName") or visit_type.get("Name", ""),
                            "status": "specialty_override_error", "message": str(override_error), "slot_count": 0,
                            "decision_tree": lqf.get("prompts", []),
                        })
                        print(f"    {summaries[-1]['appointment_type']}: skipped ({override_error}); continuing", flush=True)
                        continue
                    evaluated_visit_id = evaluated.get("VisitTypeId")
                    active_visit = next((item for item in detail.get("VisitTypes", []) if epic_id(item) == evaluated_visit_id), visit_type)
                pair_detail = detail if detail.get("ProviderDepartmentPairs") else base_detail
                pair_visit = active_visit if pair_detail is detail else visit_type
                pairs = [pair for pair in pair_detail.get("ProviderDepartmentPairs", []) if any(
                    info.get("VisitTypeID") == epic_id(pair_visit) for info in pair.get("VisitTypeInformation", []))]
                if not pairs:
                    pairs = pair_detail.get("ProviderDepartmentPairs", [])
                provider_filter = set(evaluated.get("ProvidersToSelect") or []) if evaluated.get("ReplacedAllOriginalProviders") else set()
                if provider_filter:
                    pairs = [pair for pair in pairs if pair.get("ProviderId") in provider_filter]
                request_model = build_slot_request(workflow_settings, workflow, specialty, reason, pair_visit, pairs, lqf)
                providers = {item.get("ID"): item for item in pair_detail.get("Providers", [])}
                departments = {item.get("ID"): item for item in pair_detail.get("Departments", [])}
                flow_rows = []
                error_code = ""
                slot_response_error = ""
                retry_count = 0
                loads = 0
                page_guard_reached = False
                for load_number in range(1, args.max_slot_loads + 1):
                    data = None
                    for attempt in range(args.api_retries + 1):
                        fields = postify(request_model)
                        fields.append(("__RequestVerificationToken", client.csrf_token))
                        try:
                            data = client.post_json("slots", fields, {
                                "phase": "slots", "specialty_name": name,
                                "appointment_type": pair_visit.get("DisplayName") or pair_visit.get("Name", ""),
                                "load_number": load_number, "attempt": attempt + 1,
                            })
                            slot_response_error = ""
                            break
                        except Exception as load_error:
                            slot_response_error = str(load_error)
                            if attempt >= args.api_retries:
                                break
                            retry_count += 1
                            print(
                                f"      load {load_number} returned invalid data; "
                                f"refreshing the anonymous session and retrying "
                                f"({attempt + 1}/{args.api_retries})",
                                flush=True,
                            )
                            try:
                                client.bootstrap()
                            except Exception:
                                pass
                            time.sleep(max(args.api_delay, 0.5) * (attempt + 1))
                    if data is None:
                        print(
                            f"    {pair_visit.get('DisplayName') or pair_visit.get('Name', '')}: "
                            f"slot response skipped after {args.api_retries + 1} attempt(s); continuing",
                            flush=True,
                        )
                        break
                    loads = load_number
                    error_code = data.get("ErrorCode") or ""
                    for solution in data.get("Solutions", []):
                        for slot in solution.get("Slots", []):
                            flow_rows.append(normalize_api_slot(slot, specialty, pair_visit, reason,
                                                                lqf.get("prompts", []), providers, departments, load_number))
                    request_model["continueInfo"] = data.get("ContinueInfo")
                    if error_code or not data.get("ContinueInfo") or data["ContinueInfo"].get("IsStopSearch"):
                        break
                    if load_number == 1 or load_number % args.slot_progress_every == 0:
                        print(
                            f"      {pair_visit.get('DisplayName') or pair_visit.get('Name', '')}: "
                            f"page {load_number}; {len(flow_rows)} slots so far",
                            flush=True,
                        )
                    time.sleep(args.api_delay)
                else:
                    page_guard_reached = True
                rows.extend(flow_rows)
                summaries.append({
                    "specialty": name, "appointment_type": pair_visit.get("DisplayName") or pair_visit.get("Name", ""),
                    "reason_for_visit": reason.get("DisplayName") or reason.get("Title", ""),
                    "status": "slot_response_error" if slot_response_error else ("slot_lookup_error" if error_code else ("page_guard_reached" if page_guard_reached else "slots_captured")),
                    "slot_count": len(flow_rows), "provider_count": len({row["provider_id"] for row in flow_rows}),
                    "location_count": len({row["department_id"] for row in flow_rows}), "loads_completed": loads,
                    "error_code": error_code, "message": slot_response_error, "retry_count": retry_count,
                    "page_guard_reached": page_guard_reached,
                    "decision_tree": lqf.get("prompts", []),
                })
                # Persist every completed public visit flow.  A long-running
                # specialty can therefore be interrupted without losing prior
                # flow results; the final pass still refreshes the latest copies.
                write_api_slot_artifacts(
                    artifact_dir, client.captures, rows, summaries, inventory_rows,
                    publish_latest=False,
                )
                print(f"    {summaries[-1]['appointment_type']}: {len(flow_rows)} slots", flush=True)
        unique_rows, unique_inventory, walk_in_rows = write_api_slot_artifacts(
            artifact_dir, client.captures, rows, summaries, inventory_rows, publish_latest=True)
        print(f"Wrote {len(unique_rows)} unique slots to {artifact_dir / artifact_name('slots.csv')}")
        print(f"Wrote {len(unique_inventory)} doctor-location pairs to {artifact_dir / artifact_name('doctor_locations.csv')}")
        print(f"Wrote {len(walk_in_rows)} walk-in clinic locations to {artifact_dir / artifact_name('walk_in_locations.csv')}")
    except Exception:
        summaries.append({"status": "run_error", "message": traceback.format_exc().splitlines()[-1]})
        write_api_slot_artifacts(artifact_dir, client.captures, rows, summaries, inventory_rows,
                                 publish_latest=True)
        (artifact_dir / artifact_name("error.txt")).write_text(traceback.format_exc(), encoding="utf-8")
        print(f"Direct slot extraction stopped, but partial CSVs were saved. Details: {artifact_dir / artifact_name('error.txt')}")
def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return ""
    return str(value).strip()


def find_first(obj: Any, names: tuple[str, ...]) -> str:
    wanted = {n.lower() for n in names}
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key.lower().replace("_", "") in {x.replace("_", "") for x in wanted}:
                result = text_value(value)
                if result:
                    return result
        for value in obj.values():
            result = find_first(value, names)
            if result:
                return result
    elif isinstance(obj, list):
        for value in obj:
            result = find_first(value, names)
            if result:
                return result
    return ""


def records(obj: Any, url: str, source: str):
    """Yield one row per dict that looks like a provider/location/specialty."""
    if isinstance(obj, dict):
        keys = {str(k).lower() for k in obj}
        interesting = any(any(h in k for h in ("provider", "doctor", "special", "location", "address", "postal", "zip")) for k in keys)
        if interesting:
            yield {
                "specialty": find_first(obj, ("specialty", "specialtyName", "department", "serviceLine")),
                "provider": find_first(obj, ("providerName", "displayName", "physicianName", "doctorName", "name")),
                "location": find_first(obj, ("locationName", "facilityName", "siteName", "clinicName")),
                "address": find_first(obj, ("address", "addressLine1", "street")),
                "city": find_first(obj, ("city",)),
                "state": find_first(obj, ("state", "stateCode")),
                "zip": find_first(obj, ("postalCode", "zip", "zipCode")),
                "appointment_date": find_first(obj, ("appointmentDate", "date", "startDate", "dateTime")),
                "appointment_time": find_first(obj, ("appointmentTime", "time", "startTime")),
                "source_url": url,
                "source": source,
                "raw_json": json.dumps(obj, ensure_ascii=False, separators=(",", ":")),
            }
        for value in obj.values():
            yield from records(value, url, source)
    elif isinstance(obj, list):
        for value in obj:
            yield from records(value, url, source)


def visible_controls(driver):
    return driver.execute_script("""
      const seen = new Set(), found = [];
      function scan(root) {
        for (const e of root.querySelectorAll('*')) {
          if (e.shadowRoot) scan(e.shadowRoot);
          const style = getComputedStyle(e);
          const visible = style.display !== 'none' && style.visibility !== 'hidden' &&
                          !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
          const interactive = /^(INPUT|SELECT|BUTTON|A)$/.test(e.tagName) ||
            e.getAttribute('role') || e.hasAttribute('tabindex') || e.hasAttribute('onclick');
          if (visible && interactive && !seen.has(e)) {
            seen.add(e);
            found.push({tag:e.tagName, type:e.type||'', name:e.name||'', id:e.id||'',
                        text:(e.innerText||e.value||'').trim().replace(/\\s+/g,' '),
                        aria:e.getAttribute('aria-label')||'', role:e.getAttribute('role')||''});
          }
        }
      }
      scan(document);
      return found;
    """)


def click_text(driver, patterns: tuple[str, ...], exclude: tuple[str, ...] = ()) -> bool:
    """Click the first visible button/card/label whose text matches."""
    elements = driver.find_elements(By.CSS_SELECTOR, "button,[role='button'],[role='option'],label,a,input[type='button'],input[type='submit']")
    for element in elements:
        if not element.is_displayed() or not element.is_enabled():
            continue
        text = " ".join(filter(None, [element.text, element.get_attribute("value"), element.get_attribute("aria-label"), element.get_attribute("title")])).strip().lower()
        if any(word in text for word in exclude):
            continue
        if any(word in text for word in patterns):
            dispatch_mouse_click(driver, element)
            return True
    return bool(driver.execute_script("""
      const patterns = arguments[0], exclude = arguments[1];
      function scan(root) {
        for (const e of root.querySelectorAll('*')) {
          if (e.shadowRoot) { const result = scan(e.shadowRoot); if (result) return true; }
          const text = (e.textContent || e.getAttribute('aria-label') || e.title || '').trim().toLowerCase();
          const visible = !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
          const leafOrInteractive = !e.children.length || /^(BUTTON|A|INPUT|SELECT)$/.test(e.tagName) || e.getAttribute('role') || e.hasAttribute('tabindex');
          if (visible && leafOrInteractive && text && patterns.some(p => text === p || text.includes(p)) && !exclude.some(p => text.includes(p))) {
            for (const type of ['pointerdown', 'mousedown', 'pointerup', 'mouseup'])
              e.dispatchEvent(new MouseEvent(type, {bubbles:true, cancelable:true, view:window}));
            e.dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
            return true;
          }
        }
        return false;
      }
      return scan(document);
    """, [p.lower() for p in patterns], [p.lower() for p in exclude]))


def click_exact_text(driver, text: str) -> bool:
    """Click a visible interactive element whose accessible text is exact."""
    wanted = text.strip().lower()
    for element in driver.find_elements(By.CSS_SELECTOR, "button,[role='button'],[role='option'],label,a,input[type='button'],input[type='submit']"):
        if not element.is_displayed() or not element.is_enabled():
            continue
        actual = " ".join(filter(None, [element.text, element.get_attribute("value"), element.get_attribute("aria-label"), element.get_attribute("title")])).strip().lower()
        if actual != wanted:
            continue
        dispatch_mouse_click(driver, element)
        return True
    return False


def dispatch_mouse_click(driver, element) -> None:
    """Send the same bubbling mouse sequence used by the scheduling cards."""
    try:
        # A keyboard activation creates a trusted click on these Epic cards.
        driver.execute_script("arguments[0].scrollIntoView({block:'center', inline:'center'});", element)
        element.send_keys(Keys.ENTER)
        return
    except WebDriverException:
        pass
    driver.execute_script("""
      const e = arguments[0];
      e.scrollIntoView({block:'center', inline:'center'});
      for (const type of ['pointerdown', 'mousedown', 'pointerup', 'mouseup'])
        e.dispatchEvent(new MouseEvent(type, {bubbles:true, cancelable:true, view:window}));
      e.dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
    """, element)


def select_visible_option(driver, preferred: tuple[str, ...] = ("Other",)) -> bool:
    """Select a safe visible questionnaire option and notify the page."""
    return bool(driver.execute_script("""
      const preferred = arguments[0].map(x => x.toLowerCase());
      const visible = e => {
        const s = getComputedStyle(e);
        return s.display !== 'none' && s.visibility !== 'hidden' &&
          !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
      };
      for (const select of [...document.querySelectorAll('select')].filter(visible)) {
        const options = [...select.options];
        const match = options.find(o => preferred.includes((o.textContent || '').trim().toLowerCase()) && o.value !== '');
        const fallback = options.find(o => o.value !== '' && !(o.disabled));
        const option = match || fallback;
        if (!option) continue;
        select.value = option.value;
        select.dispatchEvent(new Event('input', {bubbles:true}));
        select.dispatchEvent(new Event('change', {bubbles:true}));
        return true;
      }
      return false;
    """, list(preferred)))


def choose_radio_by_label(driver, label_text: str) -> bool:
    """Activate a radio input through its associated visible label."""
    return bool(driver.execute_script("""
      const wanted = String(arguments[0]).trim().toLowerCase();
      const visible = e => {
        const s = getComputedStyle(e);
        return s.display !== 'none' && s.visibility !== 'hidden' &&
          !!(e.offsetWidth || e.offsetHeight || e.getClientRects().length);
      };
      for (const label of [...document.querySelectorAll('label')].filter(visible)) {
        if ((label.textContent || '').trim().replace(/\\s+/g, ' ').toLowerCase() !== wanted) continue;
        const id = label.getAttribute('for');
        const input = id ? document.getElementById(id) : label.querySelector('input');
        if (!input) continue;
        label.scrollIntoView({block:'center'});
        label.click();
        input.dispatchEvent(new Event('change', {bubbles:true}));
        return true;
      }
      return false;
    """, label_text))


def click_specialty_card(driver, specialty_id: Any, name: str) -> bool:
    """Trigger the site's real specialty-card event with a mouse-like click."""
    # These cards are ordinary anchors with href="#".  Dispatch the full
    # pointer/click sequence because the scheduling page uses delegated
    # handlers and a plain WebElement.click() can be ignored while the page
    # is still finishing its event binding.
    result = driver.execute_script("""
      const wantedId = String(arguments[0] ?? '');
      const wantedName = String(arguments[1] ?? '').trim().toLowerCase();
      const cards = [...document.querySelectorAll('a[data-type="specialty"]')];
      const card = cards.find(a =>
        (wantedId && String(a.getAttribute('data-model-id') || '') === wantedId) ||
        ((a.textContent || '').trim().replace(/\\s+/g, ' ').toLowerCase() === wantedName)
      );
      if (!card) return {found:false, count:cards.length};
      card.scrollIntoView({block:'center', inline:'center'});
      for (const type of ['pointerdown', 'mousedown', 'pointerup', 'mouseup']) {
        card.dispatchEvent(new MouseEvent(type, {bubbles:true, cancelable:true, view:window}));
      }
      card.dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
      return {found:true, href:card.getAttribute('href'), modelId:card.getAttribute('data-model-id')};
    """, specialty_id, name)
    if isinstance(result, dict) and result.get("found"):
        return True
    return click_text(driver, (name,))


def collect_dom_slots(driver, source="dom"):
    """Collect appointment-like buttons and their nearby date/provider text."""
    return driver.execute_script("""
      const buttons = [];
      function scan(root) {
        for (const e of root.querySelectorAll('*')) {
          if (e.shadowRoot) scan(e.shadowRoot);
          if (e.tagName === 'BUTTON' && /(?:AM|PM|\\d{1,2}:\\d{2})/.test((e.innerText||'').trim())) buttons.push(e);
        }
      }
      scan(document);
      return buttons.map(b => {
        let node = b; let bits = [];
        for (let i=0; i<5 && node; i++, node=node.parentElement)
          bits.push((node.innerText||'').trim().replace(/\\s+/g,' '));
        return {time:(b.innerText||'').trim(), context:bits.join(' | ')};
      });
    """)


def save_network_data(driver, captured, network_state):
    """Continuously drain performance events and retain raw XHR/fetch bodies."""
    try:
        entries = driver.get_log("performance")
    except WebDriverException:
        entries = []
    for entry in entries:
        try:
            message = json.loads(entry["message"])["message"]
            method = message.get("method")
            params = message.get("params", {})
            request_id = params.get("requestId")
            if not request_id:
                continue
            if method == "Network.requestWillBeSent":
                request = params.get("request", {})
                network_state["requests"][request_id] = {
                    "method": request.get("method", ""),
                    "url": request.get("url", ""),
                    "post_data": request.get("postData"),
                    "document_url": params.get("documentURL", ""),
                    "timestamp": params.get("timestamp"),
                }
            elif method == "Network.responseReceived":
                response = params.get("response", {})
                url = response.get("url", "")
                mime = response.get("mimeType", "")
                resource_type = params.get("type", "")
                relevant = resource_type in ("XHR", "Fetch") or "json" in mime.lower() or any(h in url.lower() for h in JSON_HINTS)
                if relevant:
                    network_state["pending"][request_id] = {
                        "request_id": request_id,
                        "resource_type": resource_type,
                        "url": url,
                        "status": response.get("status"),
                        "status_text": response.get("statusText", ""),
                        "mime_type": mime,
                        "protocol": response.get("protocol", ""),
                        "request": network_state["requests"].get(request_id, {}),
                    }
            elif method == "Network.loadingFinished":
                if request_id in network_state["pending"]:
                    item = network_state["pending"].pop(request_id)
                    try:
                        result = driver.execute_cdp_cmd("Network.getResponseBody", {"requestId": request_id})
                        body = result.get("body", "")
                        item["body"] = body
                        item["base64_encoded"] = bool(result.get("base64Encoded"))
                        try:
                            item["data"] = json.loads(body) if not item["base64_encoded"] else None
                        except (TypeError, ValueError):
                            item["data"] = None
                    except WebDriverException as error:
                        item["body_error"] = str(error)
                    captured.append(item)
                network_state["requests"].pop(request_id, None)
        except (KeyError, TypeError, ValueError, WebDriverException):
            continue

    # EdgeDriver versions may not expose the performance log API. The page
    # hook below provides a browser-native fallback for XHR/fetch responses.
    try:
        hooked = driver.execute_script("return (window.__ohNetwork || []).splice(0);") or []
        for item in hooked:
            body = item.get("body", "")
            parsed = None
            try:
                parsed = json.loads(body)
            except (TypeError, ValueError):
                pass
            item["data"] = parsed
            item["source"] = "page_network_hook"
            captured.append(item)
    except WebDriverException:
        pass


NETWORK_HOOK = r"""
(function () {
  if (window.__ohNetworkInstalled) return;
  window.__ohNetworkInstalled = true;
  window.__ohNetwork = [];
  let seq = 0;
  const record = item => {
    item.id = ++seq;
    item.captured_at = new Date().toISOString();
    window.__ohNetwork.push(item);
    if (window.__ohNetwork.length > 2000) window.__ohNetwork.shift();
  };
  const originalFetch = window.fetch;
  window.fetch = async function (...args) {
    const request = args[0];
    const init = args[1] || {};
    const response = await originalFetch.apply(this, args);
    try {
      const body = await response.clone().text();
      record({type:'fetch', method:(init.method || (request && request.method) || 'GET'),
              url:response.url, status:response.status, mime_type:response.headers.get('content-type') || '',
              post_data:typeof init.body === 'string' ? init.body : null, body:body});
    } catch (e) {}
    return response;
  };
  const originalOpen = XMLHttpRequest.prototype.open;
  const originalSend = XMLHttpRequest.prototype.send;
  XMLHttpRequest.prototype.open = function (method, url) {
    this.__ohMethod = method; this.__ohUrl = url;
    return originalOpen.apply(this, arguments);
  };
  XMLHttpRequest.prototype.send = function (body) {
    this.addEventListener('load', () => {
      try {
        let text = '';
        if (typeof this.response === 'string') text = this.response;
        else if (this.response != null) text = JSON.stringify(this.response);
        record({type:'xhr', method:this.__ohMethod || 'GET', url:this.responseURL || this.__ohUrl || '',
                status:this.status, mime_type:this.getResponseHeader('content-type') || '',
                post_data:typeof body === 'string' ? body : null, body:text});
      } catch (e) {}
    });
    return originalSend.apply(this, arguments);
  };
})();
"""


def save_snapshot(driver, artifact_dir: Path, number: int):
    """Save a visual and DOM snapshot without clicking or changing the page."""
    prefix = f"{SITE_CODE}_stage_{number:03d}"
    controls = visible_controls(driver)
    (artifact_dir / f"{prefix}_controls.json").write_text(json.dumps(controls, indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / f"{prefix}_page.html").write_text(driver.page_source, encoding="utf-8")
    (artifact_dir / f"{prefix}_visible_text.txt").write_text(driver.find_element(By.TAG_NAME, "body").text, encoding="utf-8")
    (artifact_dir / f"{prefix}_meta.json").write_text(
        json.dumps({"url": driver.current_url, "title": driver.title, "captured_at": time.strftime("%Y-%m-%dT%H:%M:%S")}, indent=2),
        encoding="utf-8",
    )
    driver.save_screenshot(str(artifact_dir / f"{prefix}.png"))
    return controls


def first_dict_with_key(obj: Any, key: str):
    if isinstance(obj, dict):
        if key in obj:
            return obj
        for value in obj.values():
            found = first_dict_with_key(value, key)
            if found:
                return found
    elif isinstance(obj, list):
        for value in obj:
            found = first_dict_with_key(value, key)
            if found:
                return found
    return None


def build_workflow_catalog(captured):
    """Extract safe, useful catalog metadata while preserving the raw capture."""
    for item in captured:
        data = item.get("data")
        if not isinstance(data, dict) or not data.get("Specialties"):
            continue
        settings = data.get("WorkflowSettings", {})
        return {
            "source_url": item.get("url", ""),
            "specialties": [
                {"id": s.get("Id"), "name": s.get("Name"), "help_text": s.get("HelpText"), "photo_url": s.get("PhotoUrl")}
                for s in data.get("Specialties", [])
            ],
            "workflow": {
                "workflow_type": settings.get("WorkflowType"),
                "from_days_offset": settings.get("FromDaysOffset"),
                "to_days_offset": settings.get("ToDaysOffset"),
                "new_provider_from_days_offset": settings.get("NewProvFromDaysOffset"),
                "new_provider_to_days_offset": settings.get("NewProvToDaysOffset"),
                "show_insurance_verification_step": settings.get("ShowInsuranceVerificationStep"),
                "show_demographic_verification_step": settings.get("ShowDemographicVerificationStep"),
                "is_patient_location_step_required": settings.get("IsPatientLocationStepRequired"),
                "allow_open_scheduling_wizard": settings.get("AllowOpenSchedulingWizard"),
                "is_reservation_allowed": settings.get("IsReservationAllowed"),
                "home_organization_name": data.get("HomeOrganizationName"),
            },
        }
    return {"source_url": "", "specialties": [], "workflow": {}, "note": "No specialty catalog response found."}


def build_request_index(captured):
    """Create a compact index of endpoints and form-field names."""
    index = []
    for item in captured:
        request = item.get("request", {}) if isinstance(item.get("request"), dict) else {}
        url = item.get("url") or request.get("url", "")
        post_data = item.get("post_data") or request.get("post_data") or ""
        parsed = urlsplit(url)
        fields = sorted(parse_qs(post_data, keep_blank_values=True).keys()) if isinstance(post_data, str) else []
        data = item.get("data")
        index.append({
            "source": item.get("source", "performance_log"),
            "method": item.get("method") or request.get("method", ""),
            "resource_type": item.get("resource_type") or item.get("type", ""),
            "status": item.get("status"),
            "host": parsed.netloc,
            "path": parsed.path,
            "query_keys": sorted(parse_qs(parsed.query, keep_blank_values=True).keys()),
            "post_field_names": fields,
            "post_field_count": len(fields),
            "response_body_bytes": len(item.get("body", "") or ""),
            "response_top_level_keys": sorted(data.keys()) if isinstance(data, dict) else [],
        })
    return index


def build_entity_catalog(captured):
    """Summarize provider/location/department entities from workflow responses."""
    result = []
    seen = set()
    collections = ("Specialties", "Providers", "Departments", "Locations", "ProviderDepartmentPairs", "VisitTypes", "ReasonsForVisit", "Solutions")
    for item in captured:
        data = item.get("data")
        if not isinstance(data, dict):
            continue
        for collection in collections:
            values = data.get(collection)
            if not isinstance(values, list) or not values:
                continue
            key = (item.get("url", "").split("?", 1)[0], collection, len(values))
            if key in seen:
                continue
            seen.add(key)
            samples = []
            for value in values[:25]:
                if isinstance(value, dict):
                    sample = {k: value.get(k) for k in value if k.lower() in {
                        "id", "name", "nameutf8", "displayname", "firstname", "lastname", "title",
                        "departmentid", "providerid", "address", "addressline1", "city", "state", "postalcode", "zip", "phone",
                        "date", "startdate", "starttime", "appointmentdate", "appointmenttime", "slot", "slotuid", "solutiontype"
                    }}
                    if not sample:
                        sample = {"keys": sorted(value.keys())}
                    samples.append(sample)
            result.append({
                "source_url": item.get("url", "").split("?", 1)[0],
                "collection": collection,
                "count": len(values),
                "samples": samples,
            })
    return result


def build_slot_catalog(captured):
    """Summarize nested GetSlots results while leaving every raw slot intact."""
    result = []
    for item in captured:
        data = item.get("data")
        if not isinstance(data, dict) or not isinstance(data.get("Solutions"), list):
            continue
        total_slots = 0
        samples = []
        for solution in data["Solutions"]:
            if not isinstance(solution, dict) or not isinstance(solution.get("Slots"), list):
                continue
            total_slots += len(solution["Slots"])
            for slot in solution["Slots"][:10]:
                if isinstance(slot, dict):
                    scalar = {k: v for k, v in slot.items() if isinstance(v, (str, int, float, bool))}
                    samples.append(scalar or {"keys": sorted(slot.keys())})
                if len(samples) >= 25:
                    break
            if len(samples) >= 25:
                break
        result.append({
            "source_url": item.get("url", "").split("?", 1)[0],
            "solution_groups": len(data["Solutions"]),
            "total_nested_slots": total_slots,
            "slot_samples": samples,
            "error_code": data.get("ErrorCode"),
        })
    return result


def fetch_json_in_browser(driver, path: str, form_fields: dict[str, Any]):
    """POST same-origin form data in the live browser session."""
    return driver.execute_async_script("""
      const done = arguments[arguments.length - 1];
      const path = arguments[0], fields = arguments[1];
      const body = new URLSearchParams();
      for (const [key, value] of Object.entries(fields)) body.append(key, String(value ?? ''));
      const url = path + (path.includes('?') ? '&' : '?') + 'noCache=' + Math.random();
      fetch(url, {
        method: 'POST',
        credentials: 'same-origin',
        headers: {'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8', 'X-Requested-With': 'XMLHttpRequest'},
        body: body.toString()
      }).then(async response => {
        const text = await response.text();
        let data = null;
        try { data = JSON.parse(text); } catch (e) {}
        done({url:url, status:response.status, mime_type:response.headers.get('content-type') || '', body:text, data:data});
      }).catch(error => done({url:url, status:0, body_error:String(error), data:null}));
    """, path, form_fields)


def wait_for_scheduling_page(driver, timeout: int = 30):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            ready = driver.execute_script("return document.readyState === 'complete' && !!window.jQuery && !!window.$$WP;")
            if ready:
                return True
        except WebDriverException:
            pass
        time.sleep(0.5)
    return False


def enumerate_specialty_data(driver, captured, artifact_dir: Path, limit: int | None = None):
    """Enumerate every specialty through the site's own workflow state."""
    catalog = build_workflow_catalog(captured)
    specialties = catalog.get("specialties", [])
    if limit:
        specialties = specialties[:limit]
    endpoint = "/MyChart/Scheduling/Anonymous/GetSpecialtyData"
    print(f"Automated specialty enumeration starting for {len(specialties)} specialties.")
    for number, specialty in enumerate(specialties, start=1):
        name = specialty.get("name", "")
        if not name:
            continue
        if number > 1:
            print(f"  loading {name}", flush=True)
            try:
                driver.get(URL)
            except TimeoutException:
                print(f"  {number}/{len(specialties)} {name}: page load timed out; continuing", flush=True)
        wait_for_scheduling_page(driver)
        time.sleep(1)
        before = len(captured)
        print(f"  selecting {name}", flush=True)
        clicked = click_specialty_card(driver, specialty.get("id"), name)
        try:
            save_snapshot(driver, artifact_dir, number + 1)
            print(f"  after click URL: {driver.current_url}", flush=True)
        except WebDriverException:
            pass
        if not clicked:
            print(f"  {number}/{len(specialties)} {name}: card not found", flush=True)
            continue
        found = None
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            time.sleep(0.5)
            save_network_data(driver, captured, {"requests": {}, "pending": {}})
            for item in captured[before:]:
                if endpoint in item.get("url", "") and isinstance(item.get("data"), dict) and item["data"].get("Providers") is not None:
                    found = item
                    break
            if found:
                break
        if found:
            found["source"] = "automated_specialty_workflow"
            found["specialty_name"] = name
            print(f"  {number}/{len(specialties)} {name}: captured", flush=True)
            print(f"    advancing safe public flow for {name}", flush=True)
            if enumerate_default_appointment_flow(driver, captured, artifact_dir, name, before, number + 100):
                print(f"    {name}: appointment slots captured", flush=True)
            else:
                print(f"    {name}: no appointment-slot response reached", flush=True)
        else:
            print(f"  {number}/{len(specialties)} {name}: no valid specialty response", flush=True)
        time.sleep(0.5)
    (artifact_dir / artifact_name("specialty_enumeration.json")).write_text(
        json.dumps([item for item in captured if item.get("source") in {"automated_specialty_data", "automated_specialty_workflow"}], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def enumerate_default_appointment_flow(driver, captured, artifact_dir: Path, specialty_name: str, before: int, stage_number: int) -> bool:
    """Advance one specialty through safe public choices until GetSlots appears."""
    endpoint = "/MyChart/Scheduling/Anonymous/GetSlots"
    deadline = time.monotonic() + 90
    last_text = ""
    last_action = ""
    last_action_at = 0.0
    manual_prompts = 0
    next_manual_prompt = time.monotonic() + 30
    found_slots = False
    while time.monotonic() < deadline:
        time.sleep(0.75)
        save_network_data(driver, captured, {"requests": {}, "pending": {}})
        for item in captured[before:]:
            if endpoint not in item.get("url", ""):
                continue
            data = item.get("data")
            if isinstance(data, dict) and isinstance(data.get("Solutions"), list):
                item["source"] = "automated_slot_workflow"
                item["specialty_name"] = specialty_name
                found_slots = True
        if found_slots:
            # Allow the page to finish its first result batch before moving on.
            time.sleep(2)
            save_network_data(driver, captured, {"requests": {}, "pending": {}})
            for item in captured[before:]:
                if endpoint in item.get("url", "") and isinstance(item.get("data"), dict):
                    item["source"] = "automated_slot_workflow"
                    item["specialty_name"] = specialty_name
            save_snapshot(driver, artifact_dir, stage_number)
            return True

        if manual_prompts < 3 and time.monotonic() >= next_manual_prompt:
            try:
                page_lines = [line.strip() for line in driver.find_element(By.TAG_NAME, "body").text.splitlines() if line.strip()]
                preview = " | ".join(page_lines[:5])
            except WebDriverException:
                preview = "the current scheduling screen"
            print(f"  Edge needs help for {specialty_name}: {preview}", flush=True)
            print("  Click the next normal public option in Edge (do not choose Medicaid or emergency/911). The recorder will resume automatically; no terminal input is required.", flush=True)
            # Keep the browser-side network hook draining while a person
            # completes a stubborn step. This works even when the script is
            # launched by a Run button without an interactive stdin.
            manual_deadline = time.monotonic() + 60
            while time.monotonic() < manual_deadline:
                time.sleep(0.75)
                save_network_data(driver, captured, {"requests": {}, "pending": {}})
            manual_prompts += 1
            deadline += 90
            next_manual_prompt = time.monotonic() + 30

        try:
            visible_text = driver.find_element(By.TAG_NAME, "body").text
        except WebDriverException:
            return False
        normalized = " ".join(visible_text.lower().split())
        if normalized == last_text and time.monotonic() - last_action_at < 3:
            continue
        last_text = normalized
        action = ""
        if "what kind of appointment are you looking for" in normalized:
            action = "appointment_type"
            click_text(driver, ("Orlando Health",), exclude=("florida medical",))
        elif "are you scheduling with orlando health or florida medical clinic" in normalized:
            action = "organization"
            if not click_exact_text(driver, "Orlando Health"):
                click_text(driver, ("Orlando Health",), exclude=("florida medical",))
        elif "are you a new patient" in normalized:
            action = "new_patient"
            click_exact_text(driver, "Yes")
        elif "is the patient's age 18 or older" in normalized or "is the patient’s age 18 or older" in normalized:
            action = "adult"
            click_exact_text(driver, "Yes")
        elif "confirm that you do not have a medicaid plan" in normalized:
            action = "not_medicaid"
            choose_radio_by_label(driver, "Not Medicaid")
            time.sleep(0.5)
            click_text(driver, ("continue",), exclude=("medicaid", "911", "emergency"))
        elif "enter in your reason for visit" in normalized:
            action = "reason"
            select_visible_option(driver, ("Other",))
        elif "which locations work for you" in normalized:
            action = "any_location"
            click_exact_text(driver, "Any location")
        elif "continue" in normalized:
            action = "continue"
            click_text(driver, ("continue",), exclude=("medicaid", "911", "emergency"))
        if action and (action != last_action or time.monotonic() - last_action_at >= 3):
            if action != last_action:
                time.sleep(2)
            last_action = action
            last_action_at = time.monotonic()
            time.sleep(1.5)
    try:
        save_snapshot(driver, artifact_dir, stage_number)
    except WebDriverException:
        pass
    return False


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", help="ZIP code to enter if the page presents a ZIP field")
    parser.add_argument("--browser", choices=("chrome", "edge"), default="edge")
    parser.add_argument("--auto", action="store_true", help="Experimental broad-choice automation; guided capture is the default")
    parser.add_argument("--enumerate", action="store_true", help="Automatically collect GetSpecialtyData for every public specialty")
    parser.add_argument("--limit-specialties", type=int, help="Limit automated specialty enumeration for testing")
    parser.add_argument("--csv", help="Optional CSV output path after raw capture")
    parser.add_argument("--analyze", help="Analyze an existing network_capture.json without opening Edge")
    parser.add_argument("--api-inventory", action="store_true", help="Map public workflow/specialty APIs directly without opening Edge")
    parser.add_argument("--api-slots", action="store_true", help="Extract public appointment slots directly without opening Edge")
    parser.add_argument("--guided-browser", action="store_true", help="Open the legacy guided Edge recorder instead of direct API extraction")
    parser.add_argument("--specialty", help="Limit direct API slot extraction to a specialty name")
    parser.add_argument("--all-specialties", action="store_true", help="Ignore a site comparison allowlist and extract the full public catalog")
    parser.add_argument("--max-slot-loads", type=int, default=8, help="Maximum paginated GetSlots calls per visit flow")
    parser.add_argument("--api-retries", type=int, default=2, help="Retries for an invalid/HTML slot response before skipping that flow")
    parser.add_argument("--api-delay", type=float, default=0.15, help="Seconds between paginated public API requests")
    parser.add_argument("--slot-progress-every", type=int, default=10,
                        help="Print a progress line after this many GetSlots pages per flow")
    parser.add_argument("--artifacts", default="OH_outputs")
    parser.add_argument("--build-master", help="Build the site-labeled master workbook from an existing run folder without calling the website")
    parser.add_argument("--wait", type=int, default=12)
    args = parser.parse_args()

    if args.build_master:
        run_dir = Path(args.build_master)
        set_active_run_tag(run_dir.name)
        workbook_path = create_master_workbook(run_dir, publish_latest=True)
        print(f"Wrote {workbook_path}")
        return

    if args.analyze:
        capture_path = Path(args.analyze)
        captured = json.loads(capture_path.read_text(encoding="utf-8"))
        artifact_dir = capture_path.parent
        set_active_run_tag(artifact_dir.name)
        (artifact_dir / artifact_name("request_index.json")).write_text(json.dumps(build_request_index(captured), indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("workflow_catalog.json")).write_text(json.dumps(build_workflow_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("entity_catalog.json")).write_text(json.dumps(build_entity_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
        (artifact_dir / artifact_name("slot_catalog.json")).write_text(json.dumps(build_slot_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
        walk_in_rows = extract_walk_in_locations(captured)
        with (artifact_dir / artifact_name("walk_in_locations.csv")).open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=["clinic_name", "address", "city", "state", "zip", "hours_url", "source_url"])
            writer.writeheader()
            writer.writerows(walk_in_rows)
        print(f"Analyzed {len(captured)} captures in {artifact_dir}")
        print(f"Wrote {artifact_dir / artifact_name('request_index.json')}")
        print(f"Wrote {artifact_dir / artifact_name('workflow_catalog.json')}")
        print(f"Wrote {len(walk_in_rows)} walk-in clinic locations to {artifact_dir / artifact_name('walk_in_locations.csv')}")
        return

    if args.api_inventory:
        run_api_inventory(args)
        return

    if args.api_slots or not args.guided_browser:
        run_api_slots(args)
        return

    session_name = extraction_folder_name()
    set_active_run_tag(session_name)
    artifact_dir = Path(args.artifacts) / session_name
    artifact_dir.mkdir(parents=True, exist_ok=True)
    options = EdgeOptions() if args.browser == "edge" else ChromeOptions()
    options.page_load_strategy = "eager"
    options.add_argument("--start-maximized")
    if args.browser == "edge":
        options.set_capability("ms:loggingPrefs", {"performance": "ALL", "browser": "ALL"})
    else:
        options.set_capability("goog:loggingPrefs", {"performance": "ALL", "browser": "ALL"})
    driver = webdriver.Edge(options=options) if args.browser == "edge" else webdriver.Chrome(options=options)
    driver.set_page_load_timeout(30)
    captured: list[dict[str, Any]] = []
    network_state = {"requests": {}, "pending": {}}
    try:
        driver.execute_cdp_cmd("Network.enable", {
            "maxTotalBufferSize": 100_000_000,
            "maxResourceBufferSize": 10_000_000,
            "maxPostDataSize": 5_000_000,
        })
        try:
            driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {"source": NETWORK_HOOK})
        except WebDriverException:
            pass
        driver.get(URL)
        # Also install in the current document in case the driver loaded it
        # before the CDP new-document hook was registered.
        try:
            driver.execute_script(NETWORK_HOOK)
        except WebDriverException:
            pass
        time.sleep(args.wait)
        wait_for_scheduling_page(driver)
        if args.zip:
            for element in driver.find_elements(By.CSS_SELECTOR, "input"):
                label = " ".join(filter(None, [element.get_attribute("name"), element.get_attribute("id"), element.get_attribute("placeholder"), element.get_attribute("aria-label")])).lower()
                if "zip" in label or "postal" in label:
                    element.clear()
                    element.send_keys(args.zip)
                    break
        controls = save_snapshot(driver, artifact_dir, 1)
        save_network_data(driver, captured, network_state)
        if args.enumerate:
            enumerate_specialty_data(driver, captured, artifact_dir, args.limit_specialties)
        elif args.auto:
            # These broad choices keep the search at the aggregate level.
            # Insurance/emergency prompts are deliberately not force-clicked.
            click_text(driver, ("any location", "all locations"))
            click_text(driver, ("any provider", "all providers"))
            click_text(driver, ("continue", "search"), exclude=("medicaid", "911", "emergency"))
            wait_for_scheduling_page(driver)
            time.sleep(1)
            if any("medicaid" in (c.get("text", "").lower()) for c in visible_controls(driver)):
                print("Insurance/eligibility screen detected. Continue manually; the recorder is still running.")

        if args.enumerate:
            save_network_data(driver, captured, network_state)
        else:
            stop = threading.Event()
            def wait_for_enter():
                print("Navigate the public scheduling flow in Edge. Press Enter here when finished capturing.")
                try:
                    input()
                except EOFError:
                    print("No terminal input is available; close Edge when finished capturing.")
                else:
                    stop.set()

            threading.Thread(target=wait_for_enter, daemon=True).start()
            signature = json.dumps(controls, sort_keys=True)
            snapshot_number = 1
            last_checkpoint = time.monotonic()
            try:
                while not stop.is_set():
                    time.sleep(1)
                    save_network_data(driver, captured, network_state)
                    current_controls = visible_controls(driver)
                    current_signature = json.dumps(current_controls, sort_keys=True)
                    if current_signature != signature:
                        snapshot_number += 1
                        save_snapshot(driver, artifact_dir, snapshot_number)
                        signature = current_signature
                    if time.monotonic() - last_checkpoint >= 5:
                        (artifact_dir / artifact_name("network_capture.partial.json")).write_text(
                            json.dumps(captured, indent=2, ensure_ascii=False), encoding="utf-8"
                        )
                        last_checkpoint = time.monotonic()
                save_network_data(driver, captured, network_state)
                dom_rows = collect_dom_slots(driver)
                if dom_rows:
                    captured.append({"url": driver.current_url, "data": {"dom_slots": dom_rows}, "resource_type": "DOM"})
            except WebDriverException:
                print("Browser was closed; saving everything captured so far.")
    except Exception:
        (artifact_dir / artifact_name("error.txt")).write_text(traceback.format_exc(), encoding="utf-8")
        print(f"The run stopped with an error. Details saved to {artifact_dir / artifact_name('error.txt')}")
    finally:
        try:
            save_network_data(driver, captured, network_state)
        except WebDriverException:
            pass
        (artifact_dir / artifact_name("network_capture.partial.json")).write_text(json.dumps(captured, indent=2, ensure_ascii=False), encoding="utf-8")
        try:
            driver.quit()
        except WebDriverException:
            pass

    (artifact_dir / artifact_name("network_capture.json")).write_text(json.dumps(captured, indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / artifact_name("workflow_catalog.json")).write_text(json.dumps(build_workflow_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / artifact_name("entity_catalog.json")).write_text(json.dumps(build_entity_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
    (artifact_dir / artifact_name("slot_catalog.json")).write_text(json.dumps(build_slot_catalog(captured), indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Captured {len(captured)} raw JSON/DOM captures in {artifact_dir}")
    if args.csv:
        rows = []
        for item in captured:
            rows.extend(records(item.get("data"), item.get("url", ""), "network_json"))
            if isinstance(item.get("data"), dict) and "dom_slots" in item["data"]:
                for slot in item["data"]["dom_slots"]:
                    context = slot.get("context", "")
                    date_match = DATE_RE.search(context)
                    rows.append({"specialty": "", "provider": "", "location": "", "address": "", "city": "", "state": "", "zip": "", "appointment_date": date_match.group(0) if date_match else "", "appointment_time": slot.get("time", ""), "source_url": item.get("url", ""), "source": "dom_slot", "raw_json": json.dumps(slot, ensure_ascii=False)})
        fields = ["specialty", "provider", "location", "address", "city", "state", "zip", "appointment_date", "appointment_time", "source_url", "source", "raw_json"]
        unique = {(tuple(row.get(field, "") for field in fields)): row for row in rows}
        with open(args.csv, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(unique.values())
        print(f"Wrote {len(unique)} normalized rows to {args.csv}")


if __name__ == "__main__":
    main()
