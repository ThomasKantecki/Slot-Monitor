#!/usr/bin/env python3
"""Cardiology Access, the six-slide executive deck.

  ~/Desktop/Cardiology-Access-Deck.pptx  - the deck to present (16:9, editable)
  ~/Desktop/Cardiology-Access-Deck.html  - rehearsal twin with the same geometry
                                           (this Mac has no PowerPoint)

Every number is read from the data artifacts and the built pages at build time,
never typed. Screenshots come from assets/deck/*.png (override with
DECK_SHOT_SLOTS, DECK_SHOT_MARKETS, DECK_SHOT_INDEX); a missing screenshot stops
the build. DECK_TODAY=YYYY-MM-DD fixes the comparison window start (default: the
wall clock, like the site). DECK_FIT=warn turns fit failures into warnings.

Run from anywhere:  python3 scripts/build-exec-deck.py
"""
import base64
import html as html_mod
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from PIL import Image, ImageFont
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import deck_style as ds
from deck_style import (ROOT, DESK, CREAM, CREAM90, NAVY, MUTE, FAINT, AH, OH, WHITE, BLACK,
                        BAND_META, MONO, DISP, LOGO_RATIOS, new_presentation, box, text, base,
                        band, kicker, logo)

OUT_PPTX = DESK / "Cardiology-Access-Deck.pptx"
OUT_HTML = DESK / "Cardiology-Access-Deck.html"
SHOTS = {
    "slots": Path(os.environ.get("DECK_SHOT_SLOTS") or ROOT / "assets" / "deck" / "slot-availability.png"),
    "markets": Path(os.environ.get("DECK_SHOT_MARKETS") or ROOT / "assets" / "deck" / "market-opportunities.png"),
    "index": Path(os.environ.get("DECK_SHOT_INDEX") or ROOT / "assets" / "deck" / "provider-index.png"),
}
SHOT_PAGE = {"slots": "public/index.html", "markets": "public/market-opportunities.html", "index": "public/provider-map.html"}
SITE = "thomaskantecki.github.io/Slot-Monitor"

fmt = lambda n: f"{n:,}"


def day(iso):
    return date.fromisoformat(str(iso)[:10]).strftime("%b %-d, %Y")


def read_json(rel, optional=False):
    path = ROOT / rel
    if optional and not path.exists():
        return None
    return json.loads(path.read_text())


def embedded(page_html, element_id):
    match = re.search(rf'<script id="{element_id}" type="application/json">(.*?)</script>', page_html, re.S)
    if not match:
        raise SystemExit(f"public/provider-map.html has no <script id=\"{element_id}\">; run npm run build first")
    return json.loads(match.group(1))


# =============================================================================
# numbers
# =============================================================================
def read_numbers(today):
    manifest = read_json("data/cardiology/current/manifest.json")
    model = read_json("data/cardiology/current/slot-times-model.json")
    ah_slots, oh_slots = manifest["ah"]["physicalSlots"], manifest["oh"]["physicalSlots"]
    if (model["totals"]["ah"], model["totals"]["oh"]) != (ah_slots, oh_slots):
        raise SystemExit("slot model totals differ from the manifest; rebuild the site first")
    common_max = model["commonMaxDate"]
    if common_max != min(model["maxDateBySystem"].values()):
        raise SystemExit("commonMaxDate is not the earlier system maximum")
    if today > common_max:
        raise SystemExit(f"today ({today}) is past the comparison window ({common_max}); refresh the slot data")
    counted = defaultdict(int)
    by_zip = defaultdict(lambda: {"ah": 0, "oh": 0})
    # physicians with an in-person opening, the strictest like-for-like cut (src/slot-rules.js), for one tile
    comparable = defaultdict(int)
    physicians = {"ah": set(), "oh": set()}
    for slot in model["slots"]:
        counted[slot["y"]] += 1
        if today <= slot["d"] <= common_max:
            by_zip[model["facilities"][slot["f"]]["z"]][slot["y"]] += 1
        if slot.get("v") or model["providers"][slot["p"]].get("c") != "Physician":
            continue
        comparable[slot["y"]] += 1
        physicians[slot["y"]].add(slot["p"])
    if (counted["ah"], counted["oh"]) != (ah_slots, oh_slots):
        raise SystemExit("counted slots differ from the manifest totals")
    all_ah, all_oh = ah_slots, oh_slots  # the site shows everything published by default, so the headline is the full count
    oh_only_zips = sum(1 for v in by_zip.values() if v["oh"] > 0 and v["ah"] == 0)

    page = (ROOT / "public" / "provider-map.html").read_text()
    zdata, zroster = embedded(page, "zdata"), embedded(page, "zroster")
    card = next((s for s in zdata["specialties"] if s["name"] == "Cardiology"), None)
    if not card:
        raise SystemExit("the Provider Index page has no Cardiology specialty entry")
    distinct, mychart = {"ah": set(), "oh": set()}, {"ah": set(), "oh": set()}
    for entries in zroster.values():
        for entry in entries:
            if entry["s"] == "Cardiology":
                distinct[entry["y"]].add(entry["i"])
            if entry.get("src") == "mychart":
                mychart[entry["y"]].add(entry["i"])
    if (len(distinct["ah"]), len(distinct["oh"])) != (card["ah"], card["oh"]):
        raise SystemExit("Cardiology roster does not match the specialty totals on the page")
    byzip = read_json("data/providers-by-zip.json")
    for sys_key in ("ah", "oh"):
        if zdata["totals"][sys_key] != byzip["totals"][sys_key] + len(mychart[sys_key]):
            raise SystemExit(f"{sys_key} index total is not directory people plus MyChart additions")
    ah_capture = (read_json("data/raw/ah-directory-scrape.json", optional=True) or {}).get("fetchedAt")
    oh_capture = (read_json("data/raw/oh-directory.json", optional=True) or {}).get("fetchedAt")
    return {
        "ah_slots": ah_slots, "oh_slots": oh_slots, "read": manifest["ah"]["runId"][:10],
        "ah_max": model["maxDateBySystem"]["ah"], "oh_max": model["maxDateBySystem"]["oh"], "common_max": common_max,
        "tele": model["telemedicineSlots"], "providers": len(model["providers"]), "facilities": len(model["facilities"]),
        "all_ah": all_ah, "all_oh": all_oh, "phys_ah": len(physicians["ah"]), "phys_oh": len(physicians["oh"]),
        "oh_only_zips": oh_only_zips, "card_ah": card["ah"], "card_oh": card["oh"],
        "offices_ah": card["ahLocations"], "offices_oh": card["ohLocations"],
        "mychart_ah": len(mychart["ah"]), "mychart_oh": len(mychart["oh"]),
        "ah_capture": ah_capture, "oh_capture": oh_capture, "index_built": zdata.get("generatedAt"),
    }


# =============================================================================
# drawing
# =============================================================================
class Canvas:
    def __init__(self):
        self.prs, self.blank = new_presentation()
        self.slides = []  # recorded ops per slide, for the twin and the fit check

    def new_slide(self):
        slide = self.prs.slides.add_slide(self.blank)
        ops = []
        ds.OPS = ops
        self.slides.append(ops)
        base(slide)
        return slide

    def picture(self, slide, key, x, y, w, max_h):
        path = SHOTS[key]
        if not path.exists():
            raise SystemExit(f"missing screenshot {path}; run ~/Desktop/deck-tools/capture-slot-monitor.cjs pages")
        page = ROOT / SHOT_PAGE[key]
        if page.exists() and path.stat().st_mtime < page.stat().st_mtime:
            print(f"warning: {path.name} is older than {SHOT_PAGE[key]}; recapture before presenting", file=sys.stderr)
        with Image.open(path) as image:
            px_w, px_h = image.size
        h = w * px_h / px_w
        if h > max_h:
            h, w = max_h, max_h * px_w / px_h
        slide.shapes.add_picture(str(path), ds.Inches(x), ds.Inches(y), width=ds.Inches(w), height=ds.Inches(h))
        box(slide, x, y, w, h, fill=None, line=BLACK, line_w=2.0)
        ds.OPS.append({"kind": "pic", "path": str(path), "x": x, "y": y, "w": w, "h": h})
        return h


def bullets(slide, x, y, w, h, items, size, leading=1.3):
    paras = [[("-  ", {"font": MONO, "bold": True, "color": FAINT}), (item, {})] for item in items]
    text(slide, x, y, w, h, paras, size=size, color=MUTE, leading=leading)


def slide_title(cv, N):
    slide = cv.new_slide()
    box(slide, 0, 0, 13.34, 1.9, fill=NAVY)
    box(slide, 3.57, 2.5, 6.2, 1.1, fill=WHITE, line=BLACK, line_w=3.0)
    text(slide, 3.57, 2.76, 6.2, 0.6, [[("CARDIOLOGY ", {"color": NAVY}), ("ACCESS", {"color": AH})]],
         size=36, font=MONO, bold=True, align=PP_ALIGN.CENTER)
    ah_w, oh_w = 0.58 * LOGO_RATIOS["AH"], 0.64 * LOGO_RATIOS["OH"]
    total = ah_w + 0.3 + 0.3 + 0.3 + oh_w
    x = (13.333 - total) / 2
    logo(slide, "AH", x, 4.1, h=0.58)
    text(slide, x + ah_w + 0.3, 4.18, 0.3, 0.45, "×", size=24, font=MONO, color=FAINT, align=PP_ALIGN.CENTER)
    logo(slide, "OH", x + ah_w + 0.9, 4.07, h=0.64)
    text(slide, 0.5, 5.15, 12.33, 0.8, "Florida cardiology appointment access, AdventHealth and Orlando Health side by side",
         size=20, align=PP_ALIGN.CENTER)
    text(slide, 1.5, 6.65, 10.33, 0.4, "September 2026 · Prepared by Thomas Kantecki", size=15, font=MONO,
         color=FAINT, align=PP_ALIGN.CENTER)


def slide_what(cv, N):
    slide = cv.new_slide()
    band(slide, "WHAT IT IS", "THREE VIEWS, ONE WEBSITE")
    columns = [
        ("SLOT AVAILABILITY", [
            "Shows open cardiology appointment slots on a Florida map by ZIP or county, with a calendar and provider cards.",
            "Counts one slot per provider, location and time.",
            "Answers where each system has open appointments and how soon.",
        ]),
        ("MARKET OPPORTUNITIES", [
            "Shows the markets where Orlando Health has cardiology slots AdventHealth does not, strongest gaps first.",
            "Counts slots inside each market circle. Circles overlap, so rows do not add.",
            "Answers which markets AdventHealth should look at first.",
        ]),
        ("PROVIDER INDEX", [
            "Shows every cardiology clinician each system publishes, placed at their offices.",
            "Counts distinct people statewide and provider-office pairs by ZIP or county.",
            "Answers how many cardiologists each system fields in an area and who they are.",
        ]),
    ]
    for i, (name, items) in enumerate(columns):
        x = 0.5 + i * 4.19
        box(slide, x, 1.35, 3.95, 4.5, fill=WHITE, line=BLACK)
        box(slide, x, 1.35, 3.95, 0.62, fill=NAVY)
        text(slide, x + 0.24, 1.5, 3.6, 0.4, name, size=15, font=MONO, color=CREAM, bold=True)
        bullets(slide, x + 0.3, 2.15, 3.4, 3.6, items, size=17, leading=1.28)
    box(slide, 0.5, 6.1, 12.33, 1.0, fill=NAVY)
    text(slide, 0.8, 6.28, 11.73, 0.66,
         "All three run as a website with no logins. Every page has a data check button that shows when the data was pulled and whether the checks pass.",
         size=16, font=MONO, color=CREAM, bold=True)


def slide_sources(cv, N):
    slide = cv.new_slide()
    band(slide, "WHERE THE DATA COMES FROM", "PUBLIC PAGES ONLY")
    panels = [
        ("APPOINTMENTS", [
            "Each system's public MyChart scheduling pages for Cardiology, read without logging in.",
            "A slot is one provider, one location and one time. A slot bookable under several visit types counts once.",
            "AdventHealth also books nurse practitioners, physician assistants and video visits online. Filters narrow both sides to physicians, in-person visits or new patients.",
            f"Both systems were read on {day(N['read'])}.",
        ]),
        ("CLINICIANS", [
            "Each system's public find-a-doctor directory, one entry per clinician with their offices.",
            "Clinicians who book in MyChart but have no directory profile are added at the clinics where they take appointments.",
            "Every clinician is placed in the ZIP and county of each office.",
        ]),
    ]
    for i, (name, items) in enumerate(panels):
        x = 0.5 + i * 6.31
        box(slide, x, 1.35, 6.02, 4.6, fill=WHITE, line=BLACK)
        box(slide, x, 1.35, 6.02, 0.64, fill=NAVY)
        text(slide, x + 0.3, 1.51, 5.5, 0.4, name, size=16.5, font=MONO, color=CREAM, bold=True)
        bullets(slide, x + 0.36, 2.25, 5.35, 3.55, items, size=17, leading=1.32)
    box(slide, 0.5, 6.2, 12.33, 0.9, fill=NAVY)
    text(slide, 0.7, 6.4, 11.93, 0.5,
         "No logins, no bookings, no patient data. All of it comes from pages any patient can open.",
         size=15, font=MONO, color=CREAM, bold=True)


def slide_how(cv, N):
    slide = cv.new_slide()
    # the address keeps its case (GitHub Pages paths are case-sensitive) and is a real link
    band(slide, "HOW TO USE IT", [[(SITE, {"link": f"https://{SITE}/"})]])
    steps = [
        ("Pick a view", "The header switches between Slot Availability, Market Opportunities and Provider Index."),
        ("Set the scope", "Choose ZIP codes or counties, a center ZIP with a radius, a date window, and In-person only."),
        ("Click an area", "Click any area for its providers and open times. The i button shows the data check."),
    ]
    for i, (title, body) in enumerate(steps):
        y = 1.35 + i * 1.88
        box(slide, 0.5, y, 4.6, 1.72, fill=WHITE, line=BLACK)
        box(slide, 0.5, y, 0.62, 1.72, fill=NAVY)
        text(slide, 0.5, y + 0.56, 0.62, 0.6, str(i + 1), size=26, font=MONO, color=CREAM, bold=True, align=PP_ALIGN.CENTER)
        kicker(slide, 1.3, y + 0.2, 3.6, title, color=NAVY, size=15)
        text(slide, 1.3, y + 0.58, 3.6, 1.05, body, size=15.5, color=MUTE, leading=1.2)
    h = cv.picture(slide, "slots", 5.35, 1.35, 7.48, 4.75)
    text(slide, 5.35, 1.35 + h + 0.15, 7.48, 0.45, "Slot Availability, the landing page of the site.", size=15, color=MUTE)


def slide_maps(cv, N):
    slide = cv.new_slide()
    band(slide, "WHAT THE MAPS SHOW", "MARKET OPPORTUNITIES AND PROVIDER INDEX")
    panels = [
        ("MARKET OPPORTUNITIES", "markets", "Markets where Orlando Health has capacity AdventHealth does not, ranked with the strongest gaps first. A white core marks a ZIP with no AdventHealth slots at all."),
        ("PROVIDER INDEX", "index", "Every cardiology clinician each system publishes, plus those who only appear in MyChart, placed at their offices by ZIP or county."),
    ]
    for i, (name, key, caption) in enumerate(panels):
        x = 0.5 + i * 6.22
        box(slide, x, 1.35, 6.11, 0.5, fill=NAVY)
        text(slide, x + 0.24, 1.47, 5.6, 0.3, name, size=15, font=MONO, color=CREAM, bold=True)
        h = cv.picture(slide, key, x, 1.85, 6.11, 4.0)
        text(slide, x, 1.85 + h + 0.18, 6.11, 1.0, caption, size=15.5, color=MUTE, leading=1.25)


def slide_today(cv, N):
    slide = cv.new_slide()
    band(slide, "WHAT IT SHOWS TODAY", f"SLOTS READ {day(N['read']).upper()}")
    box(slide, 0.5, 1.35, 4.6, 5.0, fill=WHITE, line=BLACK)
    box(slide, 0.5, 1.35, 4.6, 0.55, fill=NAVY)
    text(slide, 0.72, 1.47, 4.2, 0.34, "OPEN CARDIOLOGY SLOTS", size=15, font=MONO, color=CREAM, bold=True)
    logo(slide, "AH", 0.8, 2.25, h=0.4)
    text(slide, 2.55, 2.05, 2.35, 0.7, fmt(N["ah_slots"]), size=38, font=MONO, color=AH, bold=True, align=PP_ALIGN.RIGHT)
    box(slide, 0.75, 2.95, 4.1, 0.022, fill=BLACK)
    logo(slide, "OH", 0.8, 3.15, h=0.45)
    text(slide, 2.55, 3.0, 2.35, 0.7, fmt(N["oh_slots"]), size=38, font=MONO, color=OH, bold=True, align=PP_ALIGN.RIGHT)
    box(slide, 0.75, 3.9, 4.1, 0.022, fill=BLACK)
    text(slide, 0.8, 4.1, 4.0, 2.1,
         f"Bookable cardiology appointments each system published on {day(N['read'])}. A slot is one provider, one location, one time. "
         f"AdventHealth publishes through {day(N['ah_max'])} and Orlando Health through {day(N['oh_max'])}.",
         size=15.5, color=MUTE, leading=1.25)
    tiles = [
        ("Physicians with openings", [(fmt(N["phys_ah"]), {"color": AH}), (" / ", {"color": FAINT}), (fmt(N["phys_oh"]), {"color": OH})], AH,
         "Cardiologists with at least one open in-person slot online, the same rule on both sides."),
        ("OH only ZIP codes", [(fmt(N["oh_only_zips"]), {"color": OH})], OH,
         f"ZIP codes with Orlando Health slots and no AdventHealth slots through {day(N['common_max'])}."),
        ("Cardiology clinicians", [(fmt(N["card_ah"]), {"color": AH}), (" / ", {"color": FAINT}), (fmt(N["card_oh"]), {"color": OH})], NAVY,
         f"Clinicians in the Provider Index, at {fmt(N['offices_ah'])} and {fmt(N['offices_oh'])} offices."),
        ("MyChart only clinicians", [(fmt(N["mychart_ah"]), {"color": AH}), (" / ", {"color": FAINT}), (fmt(N["mychart_oh"]), {"color": OH})], NAVY,
         "Clinicians who book in MyChart but have no directory profile."),
    ]
    for i, (label, runs, accent, caption) in enumerate(tiles):
        x, y = 5.35 + (i % 2) * 3.83, 1.35 + (i // 2) * 2.6
        box(slide, x, y, 3.65, 2.4, fill=WHITE, line=BLACK)
        box(slide, x, y, 0.18, 2.4, fill=accent)
        kicker(slide, x + 0.42, y + 0.2, 3.05, label, color=NAVY, size=15)
        text(slide, x + 0.42, y + 0.5, 3.05, 0.8, [runs], size=36, font=MONO, bold=True)
        text(slide, x + 0.42, y + 1.35, 3.05, 0.95, caption, size=15, color=MUTE, leading=1.2)
    captured = (f"Directories captured {day(N['ah_capture'])} (AdventHealth) and {day(N['oh_capture'])} (Orlando Health)."
                if N["ah_capture"] and N["oh_capture"] else f"Provider Index built {day(N['index_built'])}.")
    box(slide, 0.5, 6.55, 12.33, 0.6, fill=CREAM90, line=BLACK)
    text(slide, 0.8, 6.62, 11.73, 0.48, f"Slots read {day(N['read'])}. {captured}", size=15, color=NAVY, bold=True)


# =============================================================================
# twin and checks
# =============================================================================
FONT_FILES = {
    "native": {(DISP, False): ("/System/Library/Fonts/HelveticaNeue.ttc", 0), (DISP, True): ("/System/Library/Fonts/HelveticaNeue.ttc", 1),
               (MONO, False): ("/System/Library/Fonts/Supplemental/Courier New.ttf", 0), (MONO, True): ("/System/Library/Fonts/Supplemental/Courier New Bold.ttf", 0)},
    "substitute": {(DISP, False): ("/System/Library/Fonts/Supplemental/Arial.ttf", 0), (DISP, True): ("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 0),
                   (MONO, False): ("/System/Library/Fonts/Supplemental/Courier New.ttf", 0), (MONO, True): ("/System/Library/Fonts/Supplemental/Courier New Bold.ttf", 0)},
}
_font_cache = {}


def measure_font(family, bold, size, flavor):
    key = (family, bold, size, flavor)
    if key not in _font_cache:
        path, index = FONT_FILES[flavor][(family, bold)]
        _font_cache[key] = ImageFont.truetype(path, size=size, index=index) if Path(path).exists() else None
    return _font_cache[key]


def paragraph_lines(runs, width_pt, flavor):
    """Greedy word wrap over runs with their own fonts; returns (lines, widest_word_pt)."""
    tokens = []
    for run in runs:
        font = measure_font(run["font"], run["bold"], run["size"], flavor)
        if font is None:
            return 1, 0
        for part in re.split(r"(\s+)", run["t"]):
            if part == "":
                continue
            tokens.append((part, font.getlength(part), part.isspace()))
    lines, line_w, widest = 1, 0.0, 0.0
    for part, w, is_space in tokens:
        if not is_space:
            widest = max(widest, w)
        if line_w + w > width_pt and not is_space and line_w > 0:
            lines += 1
            line_w = w
        else:
            line_w += w
    return lines, widest


def verify(cv, N):
    problems = []
    for n, ops in enumerate(cv.slides, start=1):
        for op in ops:
            if op["kind"] != "text":
                continue
            for para in op["paras"]:
                joined = "".join(run["t"] for run in para)
                if re.search(r"[—–%]", joined) or re.search(r"TODO|TBD|lorem|placeholder|XXX", joined, re.I):
                    problems.append(f"slide {n}: copy rule broken in {joined[:50]!r}")
                for run in para:
                    if run["size"] < 15 and run["color"] != BAND_META:
                        problems.append(f"slide {n}: {run['size']}pt text {run['t'][:40]!r}")
            for flavor in ("native", "substitute"):
                total = 0.0
                for i, para in enumerate(op["paras"]):
                    lines, widest = paragraph_lines(para, op["w"] * 72, flavor)
                    size = max(run["size"] for run in para)
                    total += lines * size * op["leading"] + (size * 0.45 if i else 0)
                    if widest > op["w"] * 72:
                        problems.append(f"slide {n} ({flavor}): a word is wider than its box in {joined[:40]!r}")
                if total > op["h"] * 72 * 1.04:
                    first = "".join(run["t"] for run in op["paras"][0])[:40]
                    problems.append(f"slide {n} ({flavor}): {first!r} needs {total / 72:.2f} in, box is {op['h']:.2f} in")
        if n == 6:
            runs = " ".join(run["t"] for op in ops if op["kind"] == "text" for para in op["paras"] for run in para)
            for key in ("ah_slots", "oh_slots", "phys_ah", "phys_oh", "oh_only_zips", "card_ah", "card_oh", "offices_ah", "offices_oh", "mychart_ah", "mychart_oh"):
                if fmt(N[key]) not in runs:
                    problems.append(f"slide 6 is missing {key} = {fmt(N[key])}")
    if len(cv.prs.slides) != 6:
        problems.append(f"expected 6 slides, built {len(cv.prs.slides)}")
    return problems


def write_twin(cv, path):
    def px(v):
        return f"{v * 96:.1f}px"
    parts = []
    for ops in cv.slides:
        items = []
        for op in ops:
            if op["kind"] == "box":
                style = f"left:{px(op['x'])};top:{px(op['y'])};width:{px(op['w'])};height:{px(op['h'])};"
                style += f"background:#{op['fill']};" if op["fill"] else "background:transparent;"
                style += f"border:{op['line_w'] * 96 / 72:.1f}px solid #{op['line']};" if op["line"] else ""
                items.append(f'<div class="b" style="{style}"></div>')
            elif op["kind"] in ("logo", "pic"):
                data = base64.b64encode(Path(op["path"]).read_bytes()).decode()
                items.append(f'<img style="left:{px(op["x"])};top:{px(op["y"])};width:{px(op["w"])};height:{px(op["h"])}" src="data:image/png;base64,{data}">')
            elif op["kind"] == "text":
                align = "center" if "CENTER" in op["align"] else "right" if "RIGHT" in op["align"] else "left"
                paras = []
                for i, para in enumerate(op["paras"]):
                    spans = ""
                    for run in para:
                        span = f'<span style="font-family:\'{run["font"]}\';font-size:{run["size"] * 96 / 72:.1f}px;font-weight:{700 if run["bold"] else 400};color:#{run["color"]}">{html_mod.escape(run["t"])}</span>'
                        spans += f'<a href="{html_mod.escape(run["link"])}" style="color:inherit;text-decoration:none">{span}</a>' if run.get("link") else span
                    margin = f"margin-top:{op['size'] * 0.45 * 96 / 72:.1f}px;" if i else ""
                    paras.append(f'<p style="{margin}line-height:{op["leading"]};text-align:{align}">{spans}</p>')
                items.append(f'<div class="t" style="left:{px(op["x"])};top:{px(op["y"])};width:{px(op["w"])};height:{px(op["h"])}">{"".join(paras)}</div>')
        parts.append(f'<section class="slide">{"".join(items)}</section>')
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>Cardiology Access deck</title><style>
html,body{{margin:0;background:#222}}.slide{{position:relative;width:1280px;height:720px;margin:24px auto;overflow:hidden;background:#F5F1E8;display:none}}.slide.on{{display:block}}
.b,.t,img{{position:absolute;box-sizing:border-box}}.t p{{margin:0;white-space:pre-wrap;overflow-wrap:break-word}}img{{object-fit:fill}}
</style></head><body>{"".join(parts)}<script>
const slides=[...document.querySelectorAll('.slide')];let n=Math.max(1,Math.min(slides.length,Number(new URLSearchParams(location.search).get('slide'))||1));
const show=()=>slides.forEach((s,i)=>s.classList.toggle('on',i===n-1));show();
addEventListener('keydown',e=>{{if(e.key==='ArrowRight'&&n<slides.length)n+=1;else if(e.key==='ArrowLeft'&&n>1)n-=1;else return;show();history.replaceState(null,'','?slide='+n);}});
</script></body></html>"""
    path.write_text(doc)


def main():
    today = os.environ.get("DECK_TODAY") or date.today().isoformat()
    N = read_numbers(today)
    cv = Canvas()
    for build in (slide_title, slide_what, slide_sources, slide_how, slide_maps, slide_today):
        build(cv, N)
    problems = verify(cv, N)
    for problem in problems:
        print("fit:", problem, file=sys.stderr)
    if problems and os.environ.get("DECK_FIT") != "warn":
        raise SystemExit(f"{len(problems)} problem(s); fix them or set DECK_FIT=warn")
    cv.prs.save(str(OUT_PPTX))
    write_twin(cv, OUT_HTML)
    print(f"wrote {OUT_PPTX} and {OUT_HTML}")
    print("numbers used:", json.dumps({k: v for k, v in N.items()}, indent=None))


if __name__ == "__main__":
    main()
